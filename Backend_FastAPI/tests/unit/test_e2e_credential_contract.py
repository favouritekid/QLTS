# tests/unit/test_e2e_credential_contract.py
# -*- coding: utf-8 -*-
"""Khoá tài khoản E2E: workflow ↔ consumer ↔ workbook seed phải nói CÙNG một thứ.

Vì sao tệp này tồn tại — sự cố nightly run ``34092354121`` (SHA ``d91b8c32``):
cả **sáu** suite E2E đỏ, 0 tín hiệu test, chỉ vì ba nguồn nói ba giọng.

* Workflow chỉ khai ``TEST_USERNAME`` / ``TEST_PASSWORD``.
* Mọi spec workflow lại đọc ``E2E_ADMIN_*`` / ``E2E_OFFICER_*``.
* Thiếu biến ⇒ spec rơi về literal cứng: ``admin`` / ``Admin@12345`` (mật khẩu
  SAI — seed là ``Admin@123``) và ``vothuhien`` / ``@Matkhau123!`` (tài khoản
  KHÔNG TỒN TẠI — officer thật là ``vothithuthuhien``).

Hai chuỗi 401 ấy kích ``ACCOUNT_LOCKOUT_MAX_ATTEMPTS=5`` ⇒ khoá 15 phút ⇒ 429 ⇒
``loginViaAPI`` retry 65s × 3. Suite ``finance`` đốt 652 giây để không đo được gì.

⚠️ **TỆP NÀY KHÔNG ĐƯỢC PHÉP SKIP.** Nó là hợp đồng xác thực; nếu nguồn chuẩn
(workbook), tệp cấu hình (workflow), thư mục consumer, hay thư viện đọc workbook
biến mất, phép kiểm phải **ĐỎ**, không được biến thành ``skipped``. GitHub coi
check ``skipped`` là thành công — một hợp đồng tự tan thành skip đúng lúc nguồn
chuẩn của nó mất là fail-open, đúng lớp lỗi mà PR này sinh ra để đóng. Vì vậy ở
đây **không có** ``pytest.skip`` và **không có** ``importorskip``.

Các bất biến được tách thành ca riêng — gộp lại thì phép kiểm vẫn xanh khi vài
cái đã hỏng:

* **C1**  mọi biến credential mà consumer ĐỌC đều phải được workflow KHAI
  (trừ biến có ``test.skip`` bảo vệ — tính bằng mã, không liệt kê tay);
* **C1b** tên biến ``E2E_*`` phải theo lược đồ — bắt ca *sai tên*, thứ C1 mù;
* **C2**  workflow phải khai ĐỦ TÁM khoá, và mỗi giá trị phải BẰNG đúng hàng
  workbook, ghép theo vai;
* **C3**  literal fallback phải ghép ĐÚNG CẶP username↔password của cùng một
  hàng workbook — không chỉ "có mặt đâu đó";
* **MANAGER** phải được khai, nếu không ca IDOR skip thành xanh giả.
* **MFA runtime** phải sinh secret tạm, bootstrap đúng hai vai đặc quyền, hoàn
  tất challenge, rồi đi qua một route Casbin — ``/login`` 200 là chưa đủ.
"""
import ast
import base64
import importlib.util
import inspect
import os
import pathlib
import re
import time

import openpyxl  # KHÔNG importorskip: thiếu nó là hợp đồng hỏng, phải ĐỎ.
import pytest

pytestmark = pytest.mark.unit


def _goc_repo() -> pathlib.Path:
    for cha in pathlib.Path(__file__).resolve().parents:
        if (cha / ".github" / "workflows").is_dir():
            return cha
    raise AssertionError(
        "không tìm được gốc repo (thư mục chứa .github/workflows). Hợp đồng "
        "credential KHÔNG được skip khi thiếu — nó phải đỏ.")


GOC = _goc_repo()

DUONG_WF = GOC / ".github" / "workflows" / "nightly-regression.yml"
DUONG_MFA_GATE = GOC / ".github" / "scripts" / "nightly_mfa_gate.py"
DUONG_XLSX = GOC / "Backend_FastAPI" / "seed_data_template.xlsx"
DUONG_POLICY_TEMPLATES = (
    GOC / "Backend_FastAPI" / "app" / "casbin_config" / "policy_templates.py"
)
THU_MUC_E2E = GOC / "frontend" / "src" / "test" / "e2e"

#: ``auth.setup.ts`` nằm trong tập consumer vì project ``chromium`` khai
#: ``dependencies: ['setup']`` — nó là đường tạo ``storageState`` cho mọi spec
#: dùng project ấy. Bỏ sót nó thì hợp đồng không canh chính đường đăng nhập
#: DUY NHẤT đã hoạt động trong run 34092354121.
CONSUMER_NGOAI_SPEC = ("auth.setup.ts",)

MAU_BIEN = re.compile(
    r"\bE2E_[A-Z]+_(?:USERNAME|PASSWORD|TOTP_SECRET)\b|"
    r"\bTEST_(?:USERNAME|PASSWORD)\b"
)

#: Tám khoá BẮT BUỘC. Danh sách đóng và tường minh: nếu chỉ kiểm "khoá nào có
#: mặt thì phải đúng", xoá cả cụm ``TEST_*`` vẫn xanh.
KHOA_BAT_BUOC = (
    "E2E_ADMIN_USERNAME", "E2E_ADMIN_PASSWORD",
    "E2E_OFFICER_USERNAME", "E2E_OFFICER_PASSWORD",
    "E2E_MANAGER_USERNAME", "E2E_MANAGER_PASSWORD",
    "TEST_USERNAME", "TEST_PASSWORD",
)

#: ``(khoá username, khoá password)`` → vai trò trong cột ``vai_tro``.
#: ``TEST_*`` trỏ officer vì ``auth.setup.ts`` đăng nhập bằng tài khoản officer.
CAP_THEO_VAI = {
    ("E2E_ADMIN_USERNAME", "E2E_ADMIN_PASSWORD"): "admin",
    ("E2E_OFFICER_USERNAME", "E2E_OFFICER_PASSWORD"): "officer",
    ("E2E_MANAGER_USERNAME", "E2E_MANAGER_PASSWORD"): "manager",
    ("TEST_USERNAME", "TEST_PASSWORD"): "officer",
}

#: Cặp hằng trong consumer: hằng username ↔ hằng password của cùng một vai,
#: **và vai trò mà hàng workbook ấy PHẢI có**.
#:
#: Vai trò là phần thứ ba, không phải trang trí: một cặp username/password
#: "đăng nhập được" chưa chứng minh gì về QUYỀN. Đổi cả cặp fallback của ADMIN
#: sang một cặp officer hợp lệ thì đăng nhập vẫn 200, C3 vẫn xanh, mà spec chạy
#: kịch bản admin dưới quyền officer — hoặc 403 hàng loạt, hoặc tệ hơn: khẳng
#: định về quyền được đo trên sai chủ thể. ``CTV`` không có trong seed nên để
#: ``None`` (fallback của nó là chuỗi rỗng, đã được ``test.skip`` bảo vệ).
VAI_CUA_TIEN_TO = {"ADMIN": "admin", "OFFICER": "officer", "MANAGER": "manager",
                   "CTV": None}
TIEN_TO_HANG = tuple(VAI_CUA_TIEN_TO)

#: ``auth.setup.ts`` dùng cặp hằng không mang tiền tố vai; nó đăng nhập bằng
#: tài khoản OFFICER để dựng ``storageState``.
VAI_CUA_SETUP = "officer"

HAU_TO_HOP_LE = ("USERNAME", "PASSWORD", "TOTP_SECRET", "MFA_CODE")
BIEN_E2E_KHAC_CHO_PHEP = {"E2E_API_URL"}
VAI_HOP_LE = ("ADMIN", "OFFICER", "MANAGER", "CTV")


def _phai_ton_tai(p: pathlib.Path, mo_ta: str) -> pathlib.Path:
    assert p.exists(), (
        f"thiếu {mo_ta}: {p}\n"
        "Hợp đồng credential phải ĐỎ khi nguồn chuẩn biến mất — skip ở đây là "
        "fail-open, vì GitHub coi check `skipped` là thành công.")
    return p


class _Phien:
    """Một context ``async with AsyncSessionLocal() as <alias>:`` trong AST."""

    __slots__ = ("alias", "node", "ket_thuc")

    def __init__(self, alias: str, node: ast.AsyncWith):
        self.alias = alias
        self.node = node
        self.ket_thuc = max(
            getattr(con, "end_lineno", node.lineno) or node.lineno
            for con in ast.walk(node)
        )


TEN_HAM_BOOTSTRAP = "bootstrap_runtime_mfa"


def _la_phien(node: ast.AST) -> "_Phien | None":
    """``AsyncWith`` mở ``AsyncSessionLocal() as <tên>`` → ``_Phien``, else None."""
    if not isinstance(node, ast.AsyncWith):
        return None
    for item in node.items:
        goi = item.context_expr
        if (
            isinstance(goi, ast.Call)
            and isinstance(goi.func, ast.Name)
            and goi.func.id == "AsyncSessionLocal"
            and isinstance(item.optional_vars, ast.Name)
        ):
            return _Phien(item.optional_vars.id, node)
    return None


def _ham_bootstrap() -> ast.AsyncFunctionDef:
    """Đúng hàm ``bootstrap_runtime_mfa`` — KHÔNG quét cả module.

    Quét cả module thì một `AsyncSessionLocal` ở hàm khác cũng tính, và phép
    "đúng hai phiên" hoá ra đếm nhầm chỗ.
    """
    nguon = _phai_ton_tai(DUONG_MFA_GATE, "script cổng MFA nightly").read_text(
        encoding="utf-8"
    )
    ham = [
        n
        for n in ast.parse(nguon).body
        if isinstance(n, ast.AsyncFunctionDef) and n.name == TEN_HAM_BOOTSTRAP
    ]
    assert len(ham) == 1, "cần đúng một `async def %s`, thấy %d" % (
        TEN_HAM_BOOTSTRAP,
        len(ham),
    )
    return ham[0]


def _ham_theo_ten(ten: tuple[str, ...]) -> dict[str, ast.AST]:
    """``{tên: node}`` cho các hàm cấp module của script cổng."""
    nguon = _phai_ton_tai(DUONG_MFA_GATE, "script cổng MFA nightly").read_text(
        encoding="utf-8"
    )
    ra = {
        n.name: n
        for n in ast.parse(nguon).body
        if isinstance(n, (ast.FunctionDef, ast.AsyncFunctionDef)) and n.name in ten
    }
    thieu = sorted(set(ten) - set(ra))
    assert not thieu, "script cổng thiếu hàm: %s" % thieu
    return ra


def _hai_phien_bootstrap() -> tuple["_Phien", "_Phien"]:
    """``(phiên ghi, phiên kiểm)`` — SIBLING TRỰC TIẾP trong `try` ngoài cùng.

    Ràng buộc "sibling trực tiếp" là phần chịu lực: nếu chỉ dùng ``ast.walk``
    thì một khối verification bọc trong ``if False:`` vẫn được tính, dù nó
    KHÔNG BAO GIỜ chạy. Đứng đúng trong ``outer_try.body`` là bằng chứng nó
    nằm trên đường thi hành thật.
    """
    fn = _ham_bootstrap()

    # `try` NGOÀI THẬT = `try` cấp cao nhất chứa trực tiếp các phiên DB.
    # Hàm còn một `try` khác (kiểm khoá Fernet) nên "đúng một try" là sai mốc;
    # đó là loại mốc xanh/đỏ vì lý do không liên quan.
    ung_vien = [
        n
        for n in fn.body
        if isinstance(n, ast.Try) and any(_la_phien(c) is not None for c in n.body)
    ]
    assert len(ung_vien) == 1, (
        "cần ĐÚNG một `try` cấp cao nhất chứa các phiên DB trong %s, thấy %d"
        % (TEN_HAM_BOOTSTRAP, len(ung_vien))
    )
    outer = ung_vien[0]
    assert outer.finalbody, (
        "`try` bọc hai phiên phải có `finally` — dọn Redis phải phủ MỌI đường "
        "lỗi của cả hai phiên, kể cả đường ném ra từ phiên kiểm."
    )

    sibling = [p for p in (_la_phien(n) for n in outer.body) if p is not None]
    assert len(sibling) == 2, (
        "cần ĐÚNG hai `async with AsyncSessionLocal()` là câu lệnh CON TRỰC TIẾP "
        "của `try` ngoài; thấy %d: %s. Một khối lồng trong `if`/`for`/nhánh chết "
        "KHÔNG được tính."
        % (len(sibling), [(p.alias, p.node.lineno) for p in sibling])
    )
    ghi, kiem = sibling
    assert ghi.node.lineno < kiem.node.lineno, "phiên ghi phải đứng TRƯỚC phiên kiểm"

    # Không được có phiên thứ ba ở BẤT KỲ đâu trong hàm — kể cả nhánh chết.
    tat_ca = [p for p in (_la_phien(n) for n in ast.walk(fn)) if p is not None]
    assert len(tat_ca) == 2, (
        "hàm mở %d context AsyncSessionLocal, chỉ được đúng 2: %s"
        % (len(tat_ca), [(p.alias, p.node.lineno) for p in tat_ca])
    )
    return ghi, kiem


def _la_await_commit(node: ast.AST, alias: str) -> bool:
    """``await <alias>.commit()`` — KHÔNG đối số, KHÔNG keyword, đúng chủ thể.

    Ba dạng phải bị loại: ``db.commit`` (không gọi), ``await x.commit()`` với
    ``x`` khác, và ``await db.commit(<gì đó>)``.
    """
    if not isinstance(node, ast.Expr) or not isinstance(node.value, ast.Await):
        return False
    goi = node.value.value
    return (
        isinstance(goi, ast.Call)
        and not goi.args
        and not goi.keywords
        and isinstance(goi.func, ast.Attribute)
        and goi.func.attr == "commit"
        and isinstance(goi.func.value, ast.Name)
        and goi.func.value.id == alias
    )


def _alias_cua_execute(node: ast.AST) -> set[str]:
    """Tên biến mà ``.execute(...)`` được gọi lên, trong một cây con."""
    ra: set[str] = set()
    for n in ast.walk(node):
        if (
            isinstance(n, ast.Call)
            and isinstance(n.func, ast.Attribute)
            and n.func.attr == "execute"
            and isinstance(n.func.value, ast.Name)
        ):
            ra.add(n.func.value.id)
    return ra


def _system_roles_san_pham() -> list[str]:
    """Tên sáu vai hệ thống, đọc TỪ SẢN PHẨM bằng AST.

    Đọc bằng AST chứ không ``import``: `policy_templates.py` kéo theo cả chuỗi
    cấu hình của ứng dụng, còn hợp đồng này phải chạy được ở lát unit trần.
    Và không chép tay danh sách: chép tay thì khi sản phẩm thêm/bớt một vai,
    cổng nightly vẫn khai đúng sáu tên cũ và không ai thấy.
    """
    nguon = _phai_ton_tai(
        DUONG_POLICY_TEMPLATES, "policy_templates.py của sản phẩm"
    ).read_text(encoding="utf-8")
    for node in ast.parse(nguon).body:
        if not isinstance(node, ast.Assign):
            continue
        if not any(
            isinstance(dich, ast.Name) and dich.id == "SYSTEM_ROLES"
            for dich in node.targets
        ):
            continue
        gia_tri = ast.literal_eval(node.value)
        return [muc["name"] for muc in gia_tri]
    raise AssertionError("không tìm thấy `SYSTEM_ROLES` trong policy_templates.py")


def _ket_qua_sync_day_du(mfa_gate, **ghi_de) -> dict:
    """Map ``results`` HỢP LỆ cho cả sáu vai; ``ghi_de`` thay từng vai một.

    Mọi ca fail-closed dưới đây xuất phát từ map này và chỉ đổi ĐÚNG MỘT vai —
    nếu ca kiểm tự dựng một map thiếu vai thì nó đỏ vì lý do khác với lý do
    đang được canh.
    """
    ra = {"role:admin": {"skipped": True, "reason": mfa_gate.LY_DO_BO_QUA_ADMIN}}
    for role in mfa_gate.EXPECTED_SYNC_ROLES:
        if role != "role:admin":
            ra[role] = {"success": True}
    ra.update(ghi_de)
    return ra


@pytest.fixture(scope="module")
def mfa_gate():
    """Nạp script thật; import hỏng phải ĐỎ, không được đổi thành skip."""
    path = _phai_ton_tai(DUONG_MFA_GATE, "script cổng MFA nightly")
    spec = importlib.util.spec_from_file_location("nightly_mfa_gate", path)
    assert spec is not None and spec.loader is not None, f"không nạp được {path}"
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


@pytest.fixture(scope="module")
def bien_runtime(mfa_gate) -> set[str]:
    return set(mfa_gate.RUNTIME_ENV_KEYS)


@pytest.fixture(scope="module")
def bien_workflow_co_san(env_wf, bien_runtime) -> set[str]:
    """Biến literal ở YAML cộng biến được sinh fail-closed qua GITHUB_ENV."""
    return set(env_wf) | bien_runtime


@pytest.fixture(scope="module")
def than_wf() -> str:
    return _phai_ton_tai(DUONG_WF, "workflow nightly").read_text(encoding="utf-8")


@pytest.fixture(scope="module")
def env_wf(than_wf) -> dict[str, str]:
    """Cặp ``KEY: value`` trong khối ``env:`` cấp job của workflow nightly."""
    ra: dict[str, str] = {}
    for dong in than_wf.splitlines():
        m = re.match(r"^\s{6}([A-Z][A-Z0-9_]*):\s*(.+?)\s*$", dong)
        if m and MAU_BIEN.fullmatch(m.group(1)):
            ra[m.group(1)] = m.group(2)
    return ra


@pytest.fixture(scope="module")
def tai_khoan_seed() -> dict[str, tuple[str, str]]:
    """``{username: (mat_khau, vai_tro)}`` đọc thẳng từ workbook seed.

    Đây là NGUỒN CHUẨN: ``seed_from_xlsx.py::seed_02_tai_khoan`` băm chính cột
    ``mat_khau`` này thành ``password_hash``, nên đăng nhập thật dùng đúng nó.
    """
    ws = openpyxl.load_workbook(
        _phai_ton_tai(DUONG_XLSX, "workbook seed"), data_only=True)["2_TaiKhoan"]
    tieu_de = [str(c.value or "").strip().lower() for c in ws[1]]
    i_u, i_p, i_v = (tieu_de.index("username"), tieu_de.index("mat_khau"),
                     tieu_de.index("vai_tro"))
    ra: dict[str, tuple[str, str]] = {}
    for hang in ws.iter_rows(min_row=2, values_only=True):
        u = str(hang[i_u] or "").strip()
        if u:
            ra[u] = (str(hang[i_p] or "").strip(), str(hang[i_v] or "").strip())
    assert ra, "sheet 2_TaiKhoan rỗng — không có nguồn chuẩn để đối chiếu"
    return ra


@pytest.fixture(scope="module")
def consumer(than_wf) -> list[pathlib.Path]:
    """Tệp THẬT SỰ đăng nhập trong lượt nightly: spec trích từ ``run:`` + setup.

    Tên spec đọc từ chính dòng lệnh chạy, không liệt kê tay — danh sách gõ tay
    sẽ trôi ngay lần ai đó thêm một suite.

    ⚠️ Tệp được nhắc trong workflow mà KHÔNG tồn tại là lỗi, **không** được lọc
    im lặng: một typo tên spec sẽ làm suite ấy biến mất khỏi cả lượt chạy lẫn
    hợp đồng này, và cả hai cùng xanh.
    """
    _phai_ton_tai(THU_MUC_E2E, "thư mục spec E2E")
    ten = sorted(set(re.findall(r"src/test/e2e/([A-Za-z0-9_.-]+\.spec\.ts)", than_wf)))
    assert ten, "không trích được spec nào từ dòng `run:` của workflow"
    thieu = [t for t in ten if not (THU_MUC_E2E / t).exists()]
    assert not thieu, (
        "workflow nhắc spec KHÔNG tồn tại — suite ấy im lặng không chạy:\n"
        + "\n".join(f"    {t}" for t in thieu))
    ra = [THU_MUC_E2E / t for t in ten]
    for t in CONSUMER_NGOAI_SPEC:
        ra.append(_phai_ton_tai(THU_MUC_E2E / t, f"consumer {t}"))
    return ra


def _hang_theo_bien(than: str) -> dict[str, str]:
    """``{TÊN_BIẾN: tên hằng JS nhận nó}``."""
    return {m.group(2): m.group(1)
            for m in re.finditer(r"const\s+(\w+)\s*=\s*process\.env\.(\w+)", than)}


def _bien_bat_buoc(p: pathlib.Path) -> set[str]:
    """Biến consumer đọc mà KHÔNG có ``test.skip`` bảo vệ.

    Miễn trừ được TÍNH, không liệt kê: một biến chỉ được bỏ qua nếu chính tệp ấy
    có ``test.skip(`` nhắc tới hằng nhận nó. Gỡ ``test.skip`` sẽ tự động biến
    biến đó thành bắt buộc.
    """
    than = p.read_text(encoding="utf-8")
    cac_skip = re.findall(r"test\.skip\([^)]*\)", than)
    bat_buoc = set()
    for bien, hang in _hang_theo_bien(than).items():
        if not MAU_BIEN.fullmatch(bien):
            continue
        if not any(hang in d for d in cac_skip):
            bat_buoc.add(bien)
    return bat_buoc


class TestHopDongCredentialE2E:
    def test_c1_moi_bien_consumer_doc_deu_duoc_workflow_khai(
        self, consumer, bien_workflow_co_san
    ):
        """C1 — bao đóng TÊN biến giữa consumer và workflow."""
        thieu: dict[str, list[str]] = {}
        for p in consumer:
            for bien in sorted(_bien_bat_buoc(p) - bien_workflow_co_san):
                thieu.setdefault(bien, []).append(p.name)
        assert not thieu, (
            "consumer đọc biến mà workflow KHÔNG khai ⇒ rơi về literal cứng và "
            "đăng nhập 401:\n" + "\n".join(
                f"    {b}  ← {', '.join(f)}" for b, f in sorted(thieu.items())))

    def test_c1b_ten_bien_e2e_phai_theo_luoc_do(self, consumer):
        """C1b — tên biến ``E2E_*`` phải khớp lược đồ, không phải tên tuỳ hứng.

        Bất biến RIÊNG khỏi C1, lý do cụ thể: C1 hỏi *"biến cần có đã khai
        chưa"*, nên nó chỉ nhìn tên ĐÚNG lược đồ. Một consumer đọc
        ``E2E_ADMIN_USER`` (thiếu ``NAME``) **rơi khỏi tầm quét của C1** — đo
        thật: đột biến ấy để C1 XANH trong khi spec đăng nhập bằng chuỗi rỗng.
        """
        la = []
        for p in consumer:
            than = p.read_text(encoding="utf-8")
            for bien in sorted(set(re.findall(r"process\.env\.(E2E_\w+)", than))):
                if bien in BIEN_E2E_KHAC_CHO_PHEP:
                    continue
                m = re.fullmatch(r"E2E_([A-Z]+)_([A-Z_]+)", bien)
                if not m or m.group(1) not in VAI_HOP_LE or m.group(2) not in HAU_TO_HOP_LE:
                    la.append(f"{p.name}: {bien}")
        assert not la, (
            "consumer đọc biến E2E_* KHÔNG theo lược đồ `E2E_<VAI>_<HẬU_TỐ>`.\n"
            f"    vai hợp lệ   : {', '.join(VAI_HOP_LE)}\n"
            f"    hậu tố hợp lệ: {', '.join(HAU_TO_HOP_LE)}\n"
            "Tên lệch sẽ đọc ra chuỗi rỗng và đăng nhập bằng thông tin trống:\n"
            + "\n".join(f"    {x}" for x in la))

    def test_c2_du_tam_khoa_va_gia_tri_bang_workbook(self, env_wf, tai_khoan_seed):
        """C2 — đủ TÁM khoá, và mỗi cặp khớp đúng một hàng workbook theo vai.

        Đòi đủ tám là cố ý: nếu chỉ kiểm "khoá nào có mặt thì phải đúng", xoá cả
        cụm ``TEST_*`` (hoặc cả sáu ``E2E_*``) vẫn cho XANH.
        """
        thieu = [k for k in KHOA_BAT_BUOC if k not in env_wf]
        assert not thieu, (
            "workflow thiếu khoá credential bắt buộc: " + ", ".join(thieu))
        loi = []
        for (ku, kp), vai_can in CAP_THEO_VAI.items():
            u, p = env_wf[ku], env_wf[kp]
            if u not in tai_khoan_seed:
                loi.append(f"{ku}={u!r} không có trong workbook seed")
                continue
            mk, vai = tai_khoan_seed[u]
            if vai != vai_can:
                loi.append(f"{ku}={u!r} có vai_tro={vai!r}, cần {vai_can!r}")
            if p != mk:
                loi.append(f"{kp} lệch mật khẩu workbook của {u!r}")
        assert not loi, "workflow ↔ workbook lệch:\n" + "\n".join(f"    {x}" for x in loi)

    def test_c3_fallback_phai_ghep_dung_cap_theo_workbook(
        self, consumer, tai_khoan_seed
    ):
        """C3 — literal fallback phải là CẶP username↔password của CÙNG một hàng.

        Không đủ nếu chỉ hỏi "password này có nằm đâu đó trong workbook không":
        đổi fallback password của ADMIN sang mật khẩu của OFFICER vẫn lọt, mà
        đăng nhập thật sẽ 401. Phải ghép theo cặp hằng của cùng một vai.
        """
        loi = []
        for p in consumer:
            than = p.read_text(encoding="utf-8")
            # {tên hằng: literal fallback}
            fb = {m.group(1): m.group(2) for m in re.finditer(
                r'const\s+(\w+)\s*=\s*process\.env\.\w+\s*\|\|\s*"([^"]*)"', than)}
            for tien_to in TIEN_TO_HANG:
                ten_u = next((h for h in fb if h.endswith("USERNAME") and tien_to in h), None)
                ten_p = next((h for h in fb if h.endswith("PASSWORD") and tien_to in h), None)
                if not ten_u or not ten_p:
                    continue
                u, mk = fb[ten_u], fb[ten_p]
                if u == "" and mk == "":
                    continue
                if u not in tai_khoan_seed:
                    loi.append(f"{p.name}: {ten_u} fallback {u!r} — tài khoản KHÔNG có trong seed")
                    continue
                mk_dung, vai = tai_khoan_seed[u]
                if mk != mk_dung:
                    loi.append(
                        f"{p.name}: {ten_p} fallback {mk!r} KHÔNG phải mật khẩu của "
                        f"{u!r} (workbook: {mk_dung!r})")
                vai_can = VAI_CUA_TIEN_TO[tien_to]
                if vai_can is not None and vai != vai_can:
                    loi.append(
                        f"{p.name}: {ten_u} fallback {u!r} có vai_tro={vai!r}, cần "
                        f"{vai_can!r} — cặp đăng nhập ĐƯỢC không có nghĩa là ĐÚNG QUYỀN")
            # auth.setup.ts dùng cặp hằng không mang tiền tố vai
            if not any(t in " ".join(fb) for t in TIEN_TO_HANG):
                u = fb.get("username", "")
                mk = fb.get("password", "")
                if u or mk:
                    if u not in tai_khoan_seed:
                        loi.append(f"{p.name}: fallback username {u!r} — không có trong seed")
                    else:
                        mk_dung, vai = tai_khoan_seed[u]
                        if mk != mk_dung:
                            loi.append(
                                f"{p.name}: fallback password {mk!r} KHÔNG phải mật khẩu "
                                f"của {u!r} (workbook: {mk_dung!r})")
                        if vai != VAI_CUA_SETUP:
                            loi.append(
                                f"{p.name}: fallback {u!r} có vai_tro={vai!r}, cần "
                                f"{VAI_CUA_SETUP!r} — storageState phải dựng bằng officer")
        assert not loi, (
            "literal fallback ghép SAI CẶP — khi biến thiếu, consumer đăng nhập "
            "bằng thông tin không tồn tại và nhận 401:\n"
            + "\n".join(f"    {x}" for x in loi))

    def test_manager_duoc_khai_de_ca_idor_chay_that(self, env_wf):
        """Bất biến RIÊNG: thiếu MANAGER thì ca IDOR skip thành XANH GIẢ.

        Tách khỏi C1 có chủ ý — C1 chỉ đòi biến *bắt buộc*, mà MANAGER được
        ``test.skip`` bảo vệ nên C1 cho phép nó vắng mặt. Chính cái vắng mặt
        "hợp lệ" ấy biến một ca IDOR thật thành xanh-do-bỏ-qua.

        ⚠️ Ca này chỉ chứng minh biến ĐƯỢC KHAI. Việc ca IDOR thật sự
        *executed* (skipped=0) là phép ĐỘNG, phải đo bằng một lượt chạy.
        """
        for k in ("E2E_MANAGER_USERNAME", "E2E_MANAGER_PASSWORD"):
            assert k in env_wf, (
                f"{k} không được khai ⇒ `lead-workflow.spec.ts` gọi `test.skip` và ca "
                "*Manager IDOR list scope* biến thành xanh giả thay vì chạy thật")


class TestHopDongMfaNightly:
    def test_tap_danh_tinh_va_route_casbin_dong(self, mfa_gate):
        """Không cho cổng trôi sang sai vai hoặc một route công khai/rỗng."""
        assert mfa_gate.LOGIN_ACCOUNTS == (
            (
                "admin",
                "E2E_ADMIN_USERNAME",
                "E2E_ADMIN_PASSWORD",
                "admin",
                "E2E_ADMIN_TOTP_SECRET",
            ),
            (
                "officer",
                "E2E_OFFICER_USERNAME",
                "E2E_OFFICER_PASSWORD",
                "officer",
                None,
            ),
            (
                "manager",
                "E2E_MANAGER_USERNAME",
                "E2E_MANAGER_PASSWORD",
                "manager",
                "E2E_MANAGER_TOTP_SECRET",
            ),
            ("setup", "TEST_USERNAME", "TEST_PASSWORD", "officer", None),
        )
        assert mfa_gate.PROTECTED_ROUTE == "/api/pipeline/all"
        assert mfa_gate.CASBIN_SYNC_ROUTE == (
            "/api/admin/roles/sync-all-from-templates?dry_run=false"
        )
        assert set(mfa_gate.USER_AGENTS) == {"admin", "officer", "manager", "setup"}
        assert mfa_gate.USER_AGENTS["officer"] != mfa_gate.USER_AGENTS["setup"]

    # ------------------------------------------------------------------
    # Sinh vật liệu MFA — bốn ca dưới đây KHÔNG được để giá trị nào lọt ra.
    #
    # pytest viết lại `assert` và in NGUYÊN VĂN cả hai vế khi đỏ, kể cả chuỗi
    # `where` lồng nhau: `assert len(base64.b32decode(pairs[name])) == 20` khi
    # thất bại in ra `where 20 = len(b'...')` rồi `where b'...' = b32decode('...')`
    # — tức cả khoá thô lẫn dạng đã mã hoá. `assert set(a) == set(b)` cũng không
    # cứu được: khối `Extra items in the right set:` in đủ từng phần tử lệch.
    # Đo thật bằng pytest 8.4.2 với đúng cờ CI (`-q --tb=short`).
    #
    # Vì vậy mọi phép chạm `pairs` / giá trị / buffer đều dùng nhánh `if` +
    # `pytest.fail(..., pytrace=False)`. `pytrace=False` đặt `style="value"`
    # (`_pytest/nodes.py:418`) nên `ReprEntry` được dựng với `localsrepr=None`
    # — không có cấu trúc nào để in, kể cả khi bật `--showlocals`.
    #
    # Lớp thứ hai: `os.urandom` bị thay bằng canary CÔNG KHAI. Nếu một chuỗi
    # canary xuất hiện trong log CI thì đó là bằng chứng rò, không phải secret.
    # ------------------------------------------------------------------

    #: Canary công khai, cố ý KHÔNG ngẫu nhiên. Ba giá trị này là thứ được
    #: quét trong stdout/stderr của các lượt đỏ có chủ đích.
    CANARY_FERNET_B64 = "CANARY0MFA0ENCRYPTION0KEY0DO0NOT0USE0000000="
    CANARY_TOTP_ADMIN = "CANARYADMINTOTPSECRET22222222222"
    CANARY_TOTP_MANAGER = "CANARYMANAGERTOTPSECRET333333333"

    @classmethod
    def _va_urandom_canary(cls, monkeypatch):
        """Thay `os.urandom` bằng canary, phân phối theo KÍCH THƯỚC.

        Phân phối theo `n` chứ không theo bộ đếm: `os.urandom` là toàn cục
        trong lúc vá, nên một lời gọi lạc từ thư viện khác sẽ làm lệch thứ tự
        mà không ai thấy. Mọi `n` ngoài hợp đồng đều ném.
        """
        c32 = base64.urlsafe_b64decode(cls.CANARY_FERNET_B64)
        hang20 = [
            base64.b32decode(cls.CANARY_TOTP_ADMIN),
            base64.b32decode(cls.CANARY_TOTP_MANAGER),
        ]

        def gia(n: int) -> bytes:
            if n == 32:
                return c32
            if n == 20:
                if not hang20:
                    raise AssertionError("os.urandom(20) gọi quá 2 lần")
                return hang20.pop(0)
            raise AssertionError("os.urandom(%d) ngoài hợp đồng canary" % n)

        monkeypatch.setattr(os, "urandom", gia)

    @classmethod
    def _sinh_va_doc(cls, mfa_gate, monkeypatch, tmp_path, capfd, *, canary=True):
        """Chạy `generate_runtime_environment`, trả `(pairs, stdout, stderr)`.

        `try/finally` quanh `readouterr()` là bắt buộc: nếu hàm ném SAU khi đã
        in, pytest sẽ dump mục `---- Captured stdout call ----` nguyên vẹn —
        `pytrace=False` không chặn được mục ấy. Rút buffer ra trước thì nó rỗng.
        """
        github_env = tmp_path / "github-env"
        monkeypatch.setenv("GITHUB_ENV", str(github_env))
        if canary:
            cls._va_urandom_canary(monkeypatch)
        try:
            mfa_gate.generate_runtime_environment()
        finally:
            out, err = capfd.readouterr()
        pairs = dict(
            line.split("=", 1)
            for line in github_env.read_text(encoding="utf-8").splitlines()
        )
        return pairs, out, err

    def test_ci_khong_bat_co_lat_nguoc_lop_che(self):
        """Ba cờ pytest có thể gỡ lớp bảo vệ — CI không được bật cái nào.

        - `--fulltrace` là cờ DUY NHẤT lật được `pytrace=False`: nó ghi đè
          `style="value"` thành `"long"` (`_pytest/nodes.py:418-423`), và khi ấy
          `pytest.fail` in lại cả toán hạng lẫn locals.
        - `-l` / `--showlocals` in biến cục bộ cho mọi `assert`/`raise` còn lại.
        - `-vv` nới `saferepr`, gỡ lớp cắt chuỗi vốn đang che bớt một cách tình cờ.
        """
        goc = _goc_repo()
        nguon = {
            "pytest.ini": goc / "Backend_FastAPI" / "pytest.ini",
            "backend-test.yml": goc / ".github" / "workflows" / "backend-test.yml",
            "nightly-regression.yml": DUONG_WF,
        }
        cam = ("--fulltrace", "--showlocals", " -l ", "-vv")
        for ten, duong in nguon.items():
            if not duong.exists():
                continue
            noi_dung = duong.read_text(encoding="utf-8")
            for co in cam:
                assert co not in noi_dung, (
                    "%s bật %r — cờ này gỡ lớp che secret khi test đỏ"
                    % (ten, co.strip())
                )

    def test_generate_van_dung_csprng_cua_he_dieu_hanh(self, mfa_gate):
        """Vá canary làm mất bằng chứng "dùng CSPRNG thật" — bù bằng AST.

        Đọc trên mã nguồn (công khai, in ra vô hại): `generate_runtime_environment`
        phải gọi đúng `os.urandom(32)`, `os.urandom(20)`, `os.urandom(20)` theo
        thứ tự ấy. Đổi sang `random.randbytes` thì ca này đỏ.
        """
        fn = _ham_theo_ten(("generate_runtime_environment",))[
            "generate_runtime_environment"
        ]
        goi = [
            n
            for n in ast.walk(fn)
            if isinstance(n, ast.Call)
            and isinstance(n.func, ast.Attribute)
            and n.func.attr == "urandom"
            and isinstance(n.func.value, ast.Name)
            and n.func.value.id == "os"
        ]
        kich_thuoc = [
            g.args[0].value
            for g in sorted(goi, key=lambda g: (g.lineno, g.col_offset))
            if g.args and isinstance(g.args[0], ast.Constant)
        ]
        assert kich_thuoc == [32, 20, 20], (
            "generate phải lấy entropy từ os.urandom theo thứ tự 32/20/20, thấy %r"
            % (kich_thuoc,)
        )

    def test_sinh_ba_secret_tam_dung_hinh_dang(
        self, mfa_gate, monkeypatch, tmp_path, capfd
    ):
        """Secret phải đúng hình dạng và chỉ đi qua GITHUB_ENV."""
        monkeypatch.delenv("GITHUB_ACTIONS", raising=False)
        pairs, _out, _err = self._sinh_va_doc(mfa_gate, monkeypatch, tmp_path, capfd)

        if tuple(pairs) != mfa_gate.RUNTIME_ENV_KEYS:
            pytest.fail(
                "GITHUB_ENV phải mang đúng %r theo thứ tự, thấy %r"
                % (list(mfa_gate.RUNTIME_ENV_KEYS), sorted(pairs)),
                pytrace=False,
            )

        raw = base64.urlsafe_b64decode(pairs["MFA_ENCRYPTION_KEY"])
        if len(raw) != 32:
            pytest.fail(
                "MFA_ENCRYPTION_KEY giải ra %d byte, cần 32" % len(raw), pytrace=False
            )
        for name in ("E2E_ADMIN_TOTP_SECRET", "E2E_MANAGER_TOTP_SECRET"):
            if not re.fullmatch(r"[A-Z2-7]{32}", pairs[name]):
                pytest.fail(
                    "%s không đúng hình dạng base32 160-bit không đệm" % name,
                    pytrace=False,
                )
            if len(base64.b32decode(pairs[name])) != 20:
                pytest.fail("%s không giải ra 160 bit" % name, pytrace=False)
        if pairs["E2E_ADMIN_TOTP_SECRET"] == pairs["E2E_MANAGER_TOTP_SECRET"]:
            pytest.fail("admin và manager nhận CÙNG một TOTP secret", pytrace=False)

    def test_tren_actions_mask_dung_ba_gia_tri(
        self, mfa_gate, monkeypatch, tmp_path, capfd
    ):
        """Bên trong Actions: đúng ba dòng `::add-mask::`, đúng ba giá trị.

        Ghim `GITHUB_ACTIONS` TƯỜNG MINH thay vì dựa vào môi trường thật — nếu
        không, ca này xanh trên runner và đỏ ở máy dev (hoặc ngược lại) mà cả
        hai kết quả đều không nói gì về mã.
        """
        monkeypatch.setenv("GITHUB_ACTIONS", "true")
        pairs, out, err = self._sinh_va_doc(mfa_gate, monkeypatch, tmp_path, capfd)

        mask = [d for d in out.splitlines() if d.startswith("::add-mask::")]
        if len(mask) != 3:
            pytest.fail(
                "cần đúng 3 dòng `::add-mask::`, đếm được %d" % len(mask),
                pytrace=False,
            )

        # Quy chiếu về TÊN KHOÁ để chẩn đoán được mà không in giá trị nào.
        da_che = {d[len("::add-mask::") :] for d in mask}
        chua_che = sorted(ten for ten, v in pairs.items() if v not in da_che)
        thua = len(da_che - set(pairs.values()))
        if chua_che or thua:
            pytest.fail(
                "mask không khớp tập giá trị sinh ra: chưa che %s; %d dòng mask "
                "không ứng với giá trị nào" % (chua_che, thua),
                pytrace=False,
            )
        if err != "":
            pytest.fail("stderr phải rỗng, nhận %d ký tự" % len(err), pytrace=False)

    def test_chay_local_khong_in_secret_ra_stdout_lan_stderr(
        self, mfa_gate, monkeypatch, tmp_path, capfd
    ):
        """Ngoài Actions: KHÔNG dòng mask nào, và không giá trị nào lọt ra log.

        `::add-mask::` chỉ có nghĩa khi runner đọc được nó. Chạy tay ở máy local
        thì dòng ấy in thẳng khoá Fernet và hai TOTP secret ra terminal — và
        không có cơ chế che nào cả. Đây là ca CHỨNG MINH guard: gỡ điều kiện
        `GITHUB_ACTIONS == "true"` thì ca này đỏ, còn ca ở trên vẫn xanh.

        `delenv` là bắt buộc: trên runner GitHub biến ấy LUÔN tồn tại.
        """
        monkeypatch.delenv("GITHUB_ACTIONS", raising=False)
        pairs, out, err = self._sinh_va_doc(mfa_gate, monkeypatch, tmp_path, capfd)

        if len(pairs) != 3:
            pytest.fail(
                "secret VẪN phải được sinh và ghi vào GITHUB_ENV: đếm được %d khoá %s"
                % (len(pairs), sorted(pairs)),
                pytrace=False,
            )
        for nhan, dem in (("stdout", out), ("stderr", err)):
            if "::add-mask::" in dem:
                pytest.fail(
                    "có dòng `::add-mask::` trên %s khi chạy NGOÀI Actions" % nhan,
                    pytrace=False,
                )
            for ten, gia_tri in pairs.items():
                if gia_tri in dem:
                    pytest.fail("%s rò ra %s khi chạy local" % (ten, nhan), pytrace=False)

    @pytest.mark.parametrize("gia_tri", ["false", "", "TRUE", "1"])
    def test_chi_chuoi_true_moi_bat_mask(
        self, mfa_gate, monkeypatch, tmp_path, capfd, gia_tri
    ):
        """Chỉ đúng chuỗi `"true"` mới bật che — không kiểm truthiness.

        `GITHUB_ACTIONS=false` là "không ở trong Actions"; nếu cổng đọc bằng
        `is not None` hay truthiness thì `""` và `"false"` sẽ cho hai kết quả
        ngược nhau mà không ai giải thích được vì sao.
        """
        monkeypatch.setenv("GITHUB_ACTIONS", gia_tri)
        pairs, out, err = self._sinh_va_doc(mfa_gate, monkeypatch, tmp_path, capfd)
        for nhan, dem in (("stdout", out), ("stderr", err)):
            if "::add-mask::" in dem:
                pytest.fail(
                    "GITHUB_ACTIONS=%r vẫn bật che trên %s" % (gia_tri, nhan),
                    pytrace=False,
                )
            for ten, v in pairs.items():
                if v in dem:
                    pytest.fail(
                        "%s rò ra %s với GITHUB_ACTIONS=%r" % (ten, nhan, gia_tri),
                        pytrace=False,
                    )

    def test_khong_in_bi_mat_ra_log(self, mfa_gate):
        """Không một `print()` nào được nội suy mật khẩu / secret / mfa_token.

        Log Actions của lượt nightly là công khai với mọi người đọc được repo.
        `::add-mask::` chỉ che những giá trị ĐÃ ĐƯỢC khai báo; một `print` thêm
        vào sau đó — ví dụ khi ai đó gỡ lỗi một tài khoản đăng nhập hỏng — sẽ
        in thẳng mật khẩu seed ra log mà không có gì chặn.

        Kiểm bằng AST trên ĐÚNG các lời gọi `print`: quét chuỗi cả tệp sẽ đỏ
        oan vì mọi tên biến `password` trong thân hàm cũng khớp.
        """
        nguon = _phai_ton_tai(DUONG_MFA_GATE, "script cổng MFA nightly").read_text(
            encoding="utf-8"
        )
        cam = re.compile(r"(?i)password|secret|totp|mfa_token|encryption_key")
        vi_pham: list[tuple[int, str]] = []
        for node in ast.walk(ast.parse(nguon)):
            if not (
                isinstance(node, ast.Call)
                and isinstance(node.func, ast.Name)
                and node.func.id == "print"
            ):
                continue
            for con in ast.walk(node):
                ten = None
                if isinstance(con, ast.Name):
                    ten = con.id
                elif isinstance(con, ast.Attribute):
                    ten = con.attr
                if ten and cam.search(ten):
                    vi_pham.append((node.lineno, ten))
        assert not vi_pham, (
            "print() nội suy định danh nhạy cảm (dòng, tên): %s" % (vi_pham,)
        )

        # Đối chứng: dòng che dấu vẫn phải còn, nếu không phép cấm trên là vô
        # nghĩa — không in gì thì cũng không che gì.
        assert "::add-mask::" in nguon, "mất dòng `::add-mask::`"

    def test_workflow_phai_la_YAML_hop_le_va_khong_khoa_trung(self):
        """Hợp đồng này chưa từng PARSE workflow — chỉ dò chuỗi trên văn bản.

        Hệ quả: một `nightly-regression.yml` hỏng cú pháp YAML vẫn đi lọt toàn
        bộ bộ test, rồi nightly đỏ trên runner trong khi cổng vẫn xanh. Và khoá
        trùng trong cùng một mapping thì YAML lấy cái SAU, âm thầm vứt cái trước
        — đúng lớp lỗi mà một phép `in` trên chuỗi không bao giờ thấy.
        """
        import yaml

        class _Loader(yaml.SafeLoader):
            pass

        def _mapping(self, node, deep=False):
            thay = [self.construct_object(k, deep=True) for k, _ in node.value]
            trung = sorted({k for k in thay if thay.count(k) > 1})
            assert not trung, "khoá trùng trong cùng mapping: %s" % trung
            return yaml.SafeLoader.construct_mapping(self, node, deep)

        _Loader.add_constructor(
            yaml.resolver.BaseResolver.DEFAULT_MAPPING_TAG, _mapping
        )

        d = yaml.load(
            _phai_ton_tai(DUONG_WF, "workflow nightly").read_text(encoding="utf-8"),
            Loader=_Loader,
        )
        assert isinstance(d, dict), "workflow không parse ra mapping"
        buoc = d["jobs"]["regression"]["steps"]
        assert isinstance(buoc, list) and buoc, "job regression không có bước nào"

        # Mỗi bước CHẠY LỆNH phải có `name`: thứ tự được canh bằng tên, nên một
        # bước `run:` vô danh là lỗ hổng ngay trong phép canh thứ tự. Bước
        # `uses:` thì không cần — GitHub tự đặt tên theo action.
        vo_danh = [
            i for i, s in enumerate(buoc) if s.get("run") and not s.get("name")
        ]
        assert not vo_danh, "bước `run:` không có `name` tại chỉ số %s" % vo_danh

    def test_workflow_thu_tu_generate_env_bootstrap_preflight_suite(
        self, than_wf, mfa_gate
    ):
        """Khoá/secret phải tồn tại trước backend; bootstrap phải đứng sau seed."""
        markers = (
            "nightly_mfa_gate.py generate",
            "Generate .env.production for backend env_file load",
            "- name: Start services",
            "- name: Seed database",
            "backend python - bootstrap < .github/scripts/nightly_mfa_gate.py",
            # Sync GHI CSDL phải đứng TRƯỚC cutover: recreate trước khi CSDL
            # được ghi thì worker mới nạp đúng policy CŨ.
            "nightly_mfa_gate.py sync-casbin",
            "--force-recreate --wait --wait-timeout 120 backend",
            "- name: Chứng minh mọi worker đã nạp policy",
            "nightly_mfa_gate.py preflight",
            "- name: E2E — Lead workflow",
        )
        positions = [than_wf.find(marker) for marker in markers]
        assert all(position >= 0 for position in positions), dict(
            zip(markers, positions)
        )
        assert positions == sorted(
            positions
        ), "thứ tự generate/env/start/seed/bootstrap/preflight/suite sai"

        for name in mfa_gate.RUNTIME_ENV_KEYS:
            assert (
                f"{name}=${name}" in than_wf
            ), f".env.production không nhận cùng giá trị runtime của {name}"
        for name in (
            "E2E_ADMIN_USERNAME",
            "E2E_OFFICER_USERNAME",
            "E2E_MANAGER_USERNAME",
        ):
            assert (
                f"{name}=${name}" in than_wf
            ), f"backend bootstrap không nhận username runtime {name}"
        for name in (
            "E2E_ADMIN_PASSWORD",
            "E2E_OFFICER_PASSWORD",
            "E2E_MANAGER_PASSWORD",
            "TEST_PASSWORD",
        ):
            assert f"{name}=${name}" not in than_wf, (
                f"{name} không cần ở backend nhưng bị chép vào env container"
            )

    def test_heredoc_env_khong_chua_backtick_hay_command_substitution(self, than_wf):
        """Thân heredoc `.env.production` KHÔNG được chứa backtick hay `$(`.

        Heredoc mở bằng `<<EOF` (không trích dẫn) nên bash THỰC THI mọi
        backtick và `$(...)` trong thân nó — kể cả trong những dòng trông như
        chú thích, vì `#` không phải chú thích bên trong heredoc.

        Đo thật trên bản trước khi vá: 11 lệnh được chạy, tất cả hỏng, và các
        dòng chú thích bị cắt nát trong tệp env sinh ra. Chỉ cần một chú thích
        tương lai chứa chuỗi trùng tên một lệnh có thật là nó CHẠY.

        `$VAR` thì được giữ — chính nó là cách bước này bơm khoá runtime vào.
        """
        than = than_wf.split("cat > .env.production <<EOF", 1)
        assert len(than) == 2, "không tìm thấy heredoc .env.production"
        khoi = than[1].split("\n          EOF", 1)[0]

        bt = [d for d in khoi.split("\n") if "`" in d]
        assert not bt, "backtick trong heredoc sẽ ĐƯỢC THỰC THI: %s" % bt[:3]
        cs = [d for d in khoi.split("\n") if "$(" in d]
        assert not cs, "`$(...)` trong heredoc sẽ ĐƯỢC THỰC THI: %s" % cs[:3]

    def test_workflow_ghim_hai_worker_trong_env_file(self, than_wf):
        """`GUNICORN_WORKERS=2` phải nằm trong env_file, KHÔNG phải job env.

        Khối `environment:` của service backend (docker-compose.yml:158-187)
        không liệt kê `GUNICORN_WORKERS`, mà Compose chỉ forward những khoá có
        tên ở đó HOẶC có trong env_file. Đặt ở job env là biến chỉ sống trên
        runner — cùng lớp lỗi với CORS_ORIGINS đã ghi ngay trong tệp này.
        """
        assert "GUNICORN_WORKERS=2" in than_wf, "chưa ghim hai worker"
        assert "GUNICORN_WORKERS: " not in than_wf, (
            "khai ở khối `env:` YAML thì biến KHÔNG tới được container"
        )
        # Đúng trong heredoc `.env.production`, tức trước dấu kết `EOF`.
        heredoc = than_wf.split("cat > .env.production <<EOF", 1)
        assert len(heredoc) == 2, "không tìm thấy heredoc .env.production"
        assert "GUNICORN_WORKERS=2" in heredoc[1].split("\n          EOF", 1)[0]

    def test_workflow_khong_duoc_ha_worker_ve_mot(self, than_wf):
        """Hạ về 1 worker làm phân kỳ enforcer BIẾN MẤT khỏi tầm đo.

        Nightly sẽ xanh nhờ giấu đi chính điều kiện nó sinh ra để canh, trong
        khi production vẫn chạy 2-4 worker.
        """
        for xau in ("GUNICORN_WORKERS=1", "GUNICORN_WORKERS: 1", "GUNICORN_WORKERS=0"):
            assert xau not in than_wf, f"workflow ép {xau}"

    def test_cutover_phai_tao_container_moi(self, than_wf):
        """Cutover phải `up -d --force-recreate`, tuyệt đối không `restart`.

        Thiếu `--force-recreate`: khi model container không đổi, Compose GIỮ
        NGUYÊN container cũ và vẫn exit 0 — mọi bằng chứng phía sau đọc lại
        chính lượt boot cũ.
        `restart`: giữ nguyên container ID nên phép so ID-trước/ID-sau mất khả
        năng phân biệt, và nó không đọc lại env_file.
        """
        assert "--force-recreate" in than_wf
        assert "--no-deps" in than_wf
        assert "restart backend" not in than_wf, "dùng `restart` là sai pattern"
        assert 'NEW_ID" = "$OLD_ID' in than_wf, "không so container ID trước/sau"

    def test_cutover_phai_cho_wait(self, than_wf):
        """`--wait` là điều kiện CẦN — thiếu nó là đo một container chuyển tiếp."""
        assert "--wait --wait-timeout 120" in than_wf

    def test_cutover_bat_dung_ba_co_khoi_dong(self, than_wf):
        """Hai cờ tắt để không chạy lại migration/seed; cờ Casbin BẬT.

        `RUN_CASBIN_LOAD_ON_STARTUP=true` là thứ khiến lifespan của TỪNG worker
        gọi `enforcer.load_policy()` — đó chính là cơ chế hội tụ cả fleet.
        """
        for can in (
            "RUN_MIGRATIONS_ON_STARTUP=false",
            "RUN_SYNC_NOTIFICATION_RULES_ON_STARTUP=false",
            "RUN_CASBIN_LOAD_ON_STARTUP=true",
        ):
            assert can in than_wf, f"cutover thiếu cờ {can}"

    def test_cutover_doc_log_cua_container_MOI(self, than_wf):
        """Log phải neo theo container ID mới, và phải gộp stderr.

        gunicorn `errorlog="-"` và structlog `StreamHandler()` đều ghi ra
        STDERR: thiếu `2>&1` thì grep đọc RỖNG rồi báo nhầm "0 dòng" — một
        cổng đỏ vì lý do sai còn tệ hơn cổng không có.
        """
        assert 'docker logs "$NEW_ID" > "$LOG" 2>&1' in than_wf
        assert "docker logs backend" not in than_wf

    def test_cutover_dem_marker_phai_bang_so_worker(self, than_wf):
        """Đếm `== SO_WORKER`, không phải `>= 1`.

        Một worker nạp policy KHÔNG chứng minh worker còn lại đã nạp — đúng
        lớp lỗi 403-ngẫu-nhiên mà cả bước cutover sinh ra để đóng.
        """
        assert 'Booting worker with pid: [0-9]+' in than_wf
        # Chuỗi THUẦN ASCII: APP_ENV=test dùng JSONRenderer(ujson) nên "✅" ra
        # dạng \u2705; mẫu có emoji sẽ không khớp gì cả.
        assert (
            "grep -cF 'Casbin AsyncEnforcer initialized and policies loaded'"
            in than_wf
        )
        assert '"$BOOT" != "$SO_WORKER"' in than_wf
        assert '"$CASBIN" != "$SO_WORKER"' in than_wf

    def test_cutover_kiem_health_bang_so_sanh_BANG(self, than_wf):
        """`grep -q healthy` khớp luôn `unhealthy` — phải so BẰNG chuỗi.

        Đã có tiền lệ trong chính kho này (`scripts/deploy.sh`).
        """
        assert '"$TRANG_THAI" != "healthy"' in than_wf
        assert 'grep -q "healthy"' not in than_wf

    def test_cutover_dem_worker_song_that(self, than_wf):
        """Đếm tiến trình con của PID 1 — image không có `ps`.

        `/proc/1/task/1/children` tự loại master (master LÀ pid 1 nhờ
        `exec "$@"` ở docker-entrypoint.sh).
        """
        assert "/proc/1/task/1/children" in than_wf
        assert '"$DEM" != "$SO_WORKER"' in than_wf

    def test_cutover_bat_worker_chet_sau_boot(self, than_wf):
        """Boot đủ hai worker rồi một con chết vẫn là fleet hỏng."""
        for can in ("Worker failed to boot", "Worker exiting", "WORKER TIMEOUT"):
            assert can in than_wf, f"cutover không canh {can!r}"

    def test_workflow_tat_auto_sync_ngay_tu_start_dau(self, than_wf):
        """`AUTO_SYNC_TEMPLATES=false` loại lớp writer PER-WORKER.

        Nhờ đó lượt ĐỒNG BỘ TEMPLATE-POLICY HẬU-SEED còn đúng một writer:
        bước `sync-casbin`. Đây KHÔNG phải khẳng định `casbin_rule` chỉ có
        một writer — migration và `seed_from_xlsx` cũng ghi bảng ấy, cả hai
        đều chạy trước bước đó.

        Bật cờ này thì MỖI worker tự chạy `sync_all_roles_from_templates()`
        lúc boot (`app/main.py:463`). Với hai worker đó là hai writer cạnh
        tranh trên cùng bảng `casbin_rule`, boot song song: worker A có thể
        nạp policy xong TRƯỚC khi B ghi sync xong, rồi A giữ ảnh chụp cũ.
        Đo thật trước khi tắt: log container mới có `Drift detected in` 2 lần
        và `Auto-sync complete` 2 lần.

        Phải nằm trong env_file: khối `environment:` của service backend
        KHÔNG liệt kê `AUTO_SYNC_TEMPLATES`, nên đặt trước lệnh compose sẽ rơi
        vào hư không mà vẫn exit 0.
        """
        assert "AUTO_SYNC_TEMPLATES=false" in than_wf
        assert "AUTO_SYNC_TEMPLATES=true" not in than_wf, (
            "còn một chỗ bật auto-sync — writer auto-sync per-worker vẫn tồn tại"
        )
        assert "AUTO_SYNC_TEMPLATES: " not in than_wf, (
            "khai ở khối `env:` YAML thì biến KHÔNG tới được container"
        )
        heredoc = than_wf.split("cat > .env.production <<EOF", 1)
        assert len(heredoc) == 2, "không tìm thấy heredoc .env.production"
        assert "AUTO_SYNC_TEMPLATES=false" in heredoc[1].split("\n          EOF", 1)[0]

    def test_cutover_kiem_co_ngay_tu_container_dau_tien(self, than_wf):
        """Cờ phải đúng từ lần TẠO container đầu, không chỉ sau cutover.

        Nếu lần boot đầu vẫn bật auto-sync thì nó đã ghi CSDL trước khi bước
        `sync-casbin` chạy, và bằng chứng "0 drift" sau cutover che mất điều đó.
        """
        assert "- name: Cờ khởi động phải đúng ngay từ container đầu tiên" in than_wf
        vi_tri_start = than_wf.find("- name: Start services")
        vi_tri_kiem = than_wf.find("- name: Cờ khởi động phải đúng ngay")
        vi_tri_seed = than_wf.find("- name: Seed database")
        assert 0 < vi_tri_start < vi_tri_kiem < vi_tri_seed, (
            "bước kiểm cờ phải nằm giữa Start services và Seed database"
        )

    def test_cutover_khang_dinh_load_only_bang_khong(self, than_wf):
        """Cổng fleet phải khẳng định BẰNG 0, không so tương đối.

        Phép cũ `DRIFT == DONGBO` và `DRIFT <= SO_WORKER` đều thoả khi cả hai
        bằng 0, nên chúng XANH y hệt cho `0/0` (đúng ý đồ) lẫn `2/2` (cờ bị bật
        lại). Một cổng không phân biệt được hai ca ngược nhau thì không canh gì.
        """
        assert '"$DRIFT" != "$DONGBO"' not in than_wf, "còn phép so tương đối"
        assert '"$DRIFT" -gt "$SO_WORKER"' not in than_wf, "còn phép so tương đối"
        for bien in ("SEED", "DRIFT", "DONGBO", "KHONGDRIFT"):
            assert '"%s=$%s"' % (bien, bien) in than_wf, (
                "%s không nằm trong vòng khẳng định bằng 0" % bien
            )
        assert '[ "${CAP#*=}" != "0" ]' in than_wf, "không khẳng định BẰNG 0"

    def test_cutover_dung_ban_DAI_cua_chuoi_no_drift(self, than_wf):
        """Phải grep bản DÀI của "No drift detected".

        Bản ngắn còn xuất hiện ở `app/services/casbin_service.py` như một lý do
        skip HỢP LỆ của endpoint sync, nên đếm bản ngắn sẽ trộn hai chuyện khác
        hẳn nhau vào một con số.
        """
        assert "No drift detected - all roles match their templates" in than_wf
        assert "grep -cF 'Auto-sync complete:'" in than_wf, (
            "thiếu dấu hai chấm — mẫu lỏng hơn cần thiết"
        )

    def test_cutover_kiem_co_that_su_vao_container(self, than_wf):
        """Cờ phải được chứng minh là ĐÃ tới container, không chỉ đặt ở host.

        Và phải hỏi ĐÍCH DANH từng tên: `.Config.Env` chứa nguyên văn
        SECRET_KEY, MFA_ENCRYPTION_KEY, E2E_*_TOTP_SECRET — in cả mảng ra log
        là tự giải mật. `grep -Fx` khớp TRỌN dòng nên `=true` không dính
        `=truex`.
        """
        assert "grep -Fxq" in than_wf, "phải khớp trọn dòng khi kiểm env"
        assert "unset ENVDUMP" in than_wf, "phải xoá bản dump env sau khi dùng"
        assert "{{range .Config.Env}}" in than_wf

    def test_bootstrap_fail_closed_dung_hai_vai(self, mfa_gate):
        """Canh cấu trúc phần chạm DB; hiệu lực thật do stack cô lập đo."""
        source = DUONG_MFA_GATE.read_text(encoding="utf-8")
        privileged = {
            (account[3], account[4]) for account in mfa_gate.PRIVILEGED_ACCOUNTS
        }
        assert privileged == {
            ("admin", "E2E_ADMIN_TOTP_SECRET"),
            ("manager", "E2E_MANAGER_TOTP_SECRET"),
        }
        for required in (
            'settings.APP_ENV != "test"',
            'database_name != "qlts_test"',
            'set(settings.MFA_ENFORCE_ROLES) != {"admin", "manager"}',
            ".with_for_update()",
            "mfa_service.enable_mfa(",
            "mfa_service.decrypt_secret(",
            "await db.commit()",
            "enabled_names != set(expected_privileged)",
            "officer was changed by MFA bootstrap",
        ):
            assert required in source, f"bootstrap mất hàng rào: {required}"
        assert not re.search(
            r"\.mfa_enabled\s*=", source
        ), "bootstrap gán thẳng cờ MFA thay vì đi qua service sản phẩm"
        assert not re.search(
            r"UPDATE\s+.*mfa_enabled", source, re.I
        ), "bootstrap sửa MFA bằng SQL tay"

    def test_thieu_totp_runtime_phai_do_truoc_khi_goi_mang(
        self, mfa_gate, monkeypatch
    ):
        """Thiếu secret là lỗi cấu hình, không được rơi về đăng nhập nửa phiên."""
        values = {
            "E2E_ADMIN_USERNAME": "admin",
            "E2E_ADMIN_PASSWORD": "admin-password",
            "E2E_OFFICER_USERNAME": "officer",
            "E2E_OFFICER_PASSWORD": "officer-password",
            "E2E_MANAGER_USERNAME": "manager",
            "E2E_MANAGER_PASSWORD": "manager-password",
            "E2E_MANAGER_TOTP_SECRET": "JBSWY3DPEHPK3PXPJBSWY3DPEHPK3PXP",
            "TEST_USERNAME": "officer",
            "TEST_PASSWORD": "officer-password",
        }
        for name, value in values.items():
            monkeypatch.setenv(name, value)
        monkeypatch.setenv("E2E_API_URL", "http://localhost:8000")
        monkeypatch.delenv("E2E_ADMIN_TOTP_SECRET", raising=False)

        with pytest.raises(
            mfa_gate.GateError,
            match="missing required environment variable E2E_ADMIN_TOTP_SECRET",
        ):
            mfa_gate.preflight()

    def test_totp_preflight_dung_counter_truoc_de_khong_dau_doc_playwright(
        self, mfa_gate, monkeypatch
    ):
        """Cổng không được tiêu đúng mã hiện tại mà browser sắp dùng lại.

        Và nó phải TRẢ VỀ counter đã dùng, không để nơi khác tính lại: hai phép
        `int(time.time() // 30)` chạy cách nhau vài mili giây có thể rơi hai bên
        mốc 30 giây và cho hai số khác nhau — đúng loại lệch chỉ nổ thỉnh thoảng.
        """
        now = 1_800_000_010.0
        secret = "JBSWY3DPEHPK3PXPJBSWY3DPEHPK3PXP"
        monkeypatch.setattr(time, "time", lambda: now)

        actual, counter = mfa_gate._totp_for_preflight(secret, "test TOTP")

        assert actual == mfa_gate._totp_now(secret, "test TOTP", now=now - 30)
        assert actual != mfa_gate._totp_now(secret, "test TOTP", now=now)

        # Counter trả về phải là counter ĐÃ SINH RA chính mã ấy.
        assert counter == int(now // 30) - 1
        assert actual == mfa_gate._totp_tu_counter(secret, "test TOTP", counter)

    # ------------------------------------------------------------------
    # Căn ranh giới cửa sổ 30 giây TRƯỚC khi sinh mã.
    #
    # Mã được backend đánh giá tại thời điểm request TỚI NƠI, không phải lúc
    # sinh. Nếu độ trễ mạng vượt quá phần còn lại của cửa sổ thì counter đã
    # tụt ra ngoài `valid_window=1` và bị từ chối — mà thiết kế CẤM retry, nên
    # cổng sẽ đỏ vì một lý do chẳng liên quan tới thứ nó canh.
    #
    # Trước bản vá này, phép căn ấy KHÔNG có ca kiểm nào: gỡ hẳn bốn dòng guard
    # thì cả bộ vẫn xanh, vì ca duy nhất chạm hàm dùng `remaining = 20` nên
    # nhánh ngủ chưa từng chạy.
    # ------------------------------------------------------------------

    @staticmethod
    def _dong_ho_tien(monkeypatch, moc: float):
        """Đồng hồ giả TIẾN theo `sleep` — trả `(doc_gio, danh_sach_ngu)`.

        Khoá `time.time` bằng một hằng là bẫy: nhánh ngủ chạy xong thì
        `now = time.time()` đọc lại vẫn ra giá trị CŨ, nên counter không đổi và
        ca sẽ đỏ vì lý do sai.
        """
        dong_ho = {"t": moc}
        ngu: list[float] = []

        def _ngu(s):
            ngu.append(s)
            dong_ho["t"] += s

        monkeypatch.setattr(time, "time", lambda: dong_ho["t"])
        monkeypatch.setattr(time, "sleep", _ngu)
        return dong_ho, ngu

    def test_nguong_can_ranh_gioi_phai_lon_hon_timeout_http(self, mfa_gate):
        """Bất đẳng thức, không phải con số — nó sống sót mọi lần chỉnh số.

        ⚠️ Phạm vi: **trong mô hình giả định độ trễ từ lúc sinh mã tới lúc đánh
        giá không vượt `HTTP_TIMEOUT_SECONDS`**, điều kiện cần và đủ để mã không
        bị từ chối là `remaining > HTTP_TIMEOUT_SECONDS`. Dấu `>` là CHẶT:
        `remaining = Δ = timeout` rơi đúng mốc 30 giây và vẫn hỏng, đúng cái bẫy
        `>` / `>=` mà `_cho_counter_vuot` đã dính.

        Giả định ấy không được bảo đảm ngoài đời: timeout của `urllib` áp cho
        từng thao tác socket chứ không cho cả request, nên ngưỡng là biên vận
        hành có dự phòng, không phải deadline tuyệt đối.

        Cận trên 29,75 vì sau `sleep(remaining + 0.25)` thì remaining mới đúng
        bằng `30 - 0.25`; ngưỡng lớn hơn số đó khiến guard một-lần không đạt nổi
        hậu điều kiện của chính nó.
        """
        assert mfa_gate.TOTP_MIN_REMAINING_SECONDS > mfa_gate.HTTP_TIMEOUT_SECONDS, (
            "ngưỡng %r không lớn hơn timeout %r — còn dải remaining ∈ [ngưỡng, "
            "timeout] khiến mã tới nơi ở cửa sổ sau"
            % (mfa_gate.TOTP_MIN_REMAINING_SECONDS, mfa_gate.HTTP_TIMEOUT_SECONDS)
        )
        can_tren = mfa_gate.TOTP_STEP_SECONDS - mfa_gate.TOTP_BOUNDARY_OVERSHOOT_SECONDS
        assert mfa_gate.TOTP_MIN_REMAINING_SECONDS <= can_tren, (
            "ngưỡng %r vượt %r — guard một lần không thể đạt hậu điều kiện"
            % (mfa_gate.TOTP_MIN_REMAINING_SECONDS, can_tren)
        )
        assert (
            mfa_gate.TOTP_MIN_REMAINING_SECONDS
            == mfa_gate.HTTP_TIMEOUT_SECONDS + mfa_gate.TOTP_BOUNDARY_SAFETY_SECONDS
        ), "ngưỡng phải DẪN XUẤT từ timeout, không phải một số độc lập"

    def test_con_sau_giay_thi_phai_CHO_sang_cua_so_ke(self, mfa_gate, monkeypatch):
        """Ca biên dưới: `remaining = 6` — nằm trong dải nguy hiểm cũ [5, 10].

        Bản trước bản vá KHÔNG chờ ở đây (ngưỡng 5), nên mã của counter trước
        có thể tới backend sau mốc 30 giây và thành `current-2`.
        """
        secret = "JBSWY3DPEHPK3PXPJBSWY3DPEHPK3PXP"
        # now = 30*c + (30 - R) với c = 60_000_000, R = 6
        dong_ho, ngu = self._dong_ho_tien(monkeypatch, 1_800_000_024.0)

        _ma, counter = mfa_gate._totp_for_preflight(secret, "test TOTP")

        assert ngu == [6.0 + mfa_gate.TOTP_BOUNDARY_OVERSHOOT_SECONDS], (
            "phải ngủ đúng phần còn lại + overshoot, thấy %r" % (ngu,)
        )
        assert counter == 60_000_000, (
            "sau khi căn lại phải lấy counter của cửa sổ MỚI trừ một, thấy %d"
            % counter
        )
        con_lai = mfa_gate.TOTP_STEP_SECONDS - (dong_ho["t"] % mfa_gate.TOTP_STEP_SECONDS)
        assert con_lai > mfa_gate.HTTP_TIMEOUT_SECONDS, (
            "sau khi căn vẫn còn %r giây — chưa thoát dải nguy hiểm" % con_lai
        )

    def test_con_dung_bang_nguong_thi_KHONG_duoc_cho(self, mfa_gate, monkeypatch):
        """Ca biên trên: `remaining` đúng bằng ngưỡng ⇒ đi thẳng.

        Không có ca này thì nâng ngưỡng lên 25 cũng xanh — hai ca kẹp ngưỡng từ
        hai phía, một mình mỗi ca đều không đủ.
        """
        secret = "JBSWY3DPEHPK3PXPJBSWY3DPEHPK3PXP"
        moc = 30.0 * 60_000_000 + (
            mfa_gate.TOTP_STEP_SECONDS - mfa_gate.TOTP_MIN_REMAINING_SECONDS
        )
        _dong_ho, ngu = self._dong_ho_tien(monkeypatch, moc)

        _ma, counter = mfa_gate._totp_for_preflight(secret, "test TOTP")

        assert ngu == [], "remaining == ngưỡng thì KHÔNG được ngủ, đã ngủ %r" % (ngu,)
        assert counter == 59_999_999

    @staticmethod
    def _bi_tu_choi(remaining: float, delay: float, mfa_gate) -> bool:
        """MÔ HÌNH backend, không phải phép đo trên backend thật.

        Backend chấp nhận `c ∈ {n'-1, n', n'+1}` với `n'` tại lúc ĐÁNH GIÁ. Vì
        đồng hồ tiến, chấp nhận ⟺ `n' = n`, tức ⟺ `delay < remaining`. Mô hình
        này bỏ qua lệch đồng hồ giữa runner và container.
        """
        return delay >= remaining

    def test_mo_hinh_tu_choi_khop_voi_nguong_da_chon(self, mfa_gate):
        """Kiểm bằng số TRONG MÔ HÌNH: không Δ nào ≤ timeout giết được mã.

        "Độ trễ hợp lệ" ở đây nghĩa là `Δ ≤ HTTP_TIMEOUT_SECONDS` — biên của mô
        hình, không phải biên của thực tế. Một request chậm bất thường vẫn có
        thể vượt nó, và khi ấy kết luận này không áp dụng.

        Và chiều ngược: ở ngưỡng cũ (5) thì dải [5, 10] vẫn chết — nếu không có
        vế này thì phép trên xanh vì lý do rỗng.
        """
        buoc = mfa_gate.TOTP_STEP_SECONDS
        tmax = mfa_gate.HTTP_TIMEOUT_SECONDS

        # Với ngưỡng MỚI: mọi remaining sống sót guard đều an toàn với mọi Δ hợp lệ.
        r = mfa_gate.TOTP_MIN_REMAINING_SECONDS
        while r <= buoc:
            assert not self._bi_tu_choi(r, tmax, mfa_gate), (
                "remaining=%r vẫn bị từ chối ở Δ=%r" % (r, tmax)
            )
            r += 0.25

        # Với ngưỡng CŨ (5): tồn tại ca chết — đối chứng cho vế trên.
        chet = [
            x / 4
            for x in range(20, 41)  # 5.00 .. 10.00
            if self._bi_tu_choi(x / 4, tmax, mfa_gate)
        ]
        assert chet, "mô hình không tái hiện được dải hở cũ — phép trên vô nghĩa"
        assert min(chet) == 5.0 and max(chet) == 10.0

    @staticmethod
    def _jar_cua(mfa_gate, opener):
        """CookieJar THẬT đang gắn vào opener — nguồn duy nhất để so danh tính.

        So bằng `is` với chính jar này mới loại được đột biến "truyền một jar
        mới dựng": jar rỗng vẫn có đủ API, chỉ là không mang cookie phiên nào.
        """
        jars = [
            h.cookiejar
            for h in opener.handlers
            if isinstance(h, mfa_gate.urllib.request.HTTPCookieProcessor)
        ]
        assert len(jars) == 1, "opener phải có đúng một HTTPCookieProcessor"
        return jars[0]

    def test_preflight_mfa_xong_moi_probe_casbin(self, mfa_gate, monkeypatch):
        """Trình tự NĂM sự kiện, đo trên HÀM THẬT — không dò chuỗi nguồn.

        Giữa challenge và verify có một phép đo BẮT BUỘC: gọi route Casbin và
        đòi ĐÚNG 401. Thiếu bước ấy thì cổng không phân biệt được "mới qua mật
        khẩu" với "đã có phiên" — xem
        `test_pre_verify_route_200_la_ro_phien_phai_do`.

        Và callback sync Casbin phải nằm ĐÚNG giữa phép kiểm vai và probe cuối:
        chạy sau probe thì cái 200 kia được đo trên policy CHƯA hội tụ, tức
        cổng xanh nhờ một trạng thái nó chưa hề kiểm.
        """
        su_kien: list[tuple] = []
        openers: list[object] = []
        responses = [
            (200, {"mfa_required": True, "mfa_token": "challenge"}),
            (401, {"detail": "Not authenticated"}),   # chưa verify ⇒ chưa có phiên
            (200, {"user": {"role": "admin"}}),
            (200, []),
        ]

        def fake_request(opener, method, url, **kwargs):
            openers.append(opener)
            su_kien.append((method, url, kwargs))
            return responses.pop(0)

        monkeypatch.setattr(mfa_gate, "_request_json", fake_request)
        monkeypatch.setattr(
            mfa_gate, "_totp_for_preflight", lambda *args: ("123456", 55_000_000)
        )
        secret = "JBSWY3DPEHPK3PXPJBSWY3DPEHPK3PXP"
        phien = mfa_gate._preflight_account(
            ("admin", "admin", "password", "admin", secret),
            "http://localhost:8000",
        )

        buoc = [e[0] for e in su_kien]
        assert buoc == ["POST", "GET", "POST", "GET"], "trình tự sai: %r" % (buoc,)
        assert su_kien[0][1].endswith("/api/auth/login")
        assert su_kien[1][1].endswith(mfa_gate.PROTECTED_ROUTE)  # probe TRƯỚC verify
        assert su_kien[2][1].endswith("/api/auth/verify-mfa")
        assert su_kien[2][2]["payload"]["mfa_token"] == "challenge"
        assert re.fullmatch(r"\d{6}", su_kien[2][2]["payload"]["code"])
        assert su_kien[3][1].endswith(mfa_gate.PROTECTED_ROUTE)  # probe SAU verify

        # Phiên trả về phải mang ĐÚNG opener/jar đã dùng và counter đã tiêu —
        # đó là thứ bước `sync-casbin` ghi ra tệp và `preflight` chờ qua.
        assert openers and all(o is openers[0] for o in openers), (
            "các request không dùng chung một opener"
        )
        assert phien.opener is openers[0]
        assert phien.cookie_jar is self._jar_cua(mfa_gate, openers[0])
        assert phien.counter_da_tieu == 55_000_000
        assert phien.label == "admin"

    def test_sync_casbin_can_csrf_va_fail_closed_theo_tung_vai(
        self, mfa_gate, monkeypatch
    ):
        cookie_jar = mfa_gate.http.cookiejar.CookieJar()
        cookie_jar.set_cookie(
            mfa_gate.http.cookiejar.Cookie(
                version=0,
                name="csrf_token",
                value="csrf-value",
                port=None,
                port_specified=False,
                domain="backend.local",
                domain_specified=True,
                domain_initial_dot=False,
                path="/",
                path_specified=True,
                secure=False,
                expires=None,
                discard=True,
                comment=None,
                comment_url=None,
                rest={},
                rfc2109=False,
            )
        )
        opener = mfa_gate.urllib.request.build_opener()
        calls = []

        def fake_request(*args, **kwargs):
            calls.append((args, kwargs))
            return (
                200,
                {"dry_run": False, "results": _ket_qua_sync_day_du(mfa_gate)},
            )

        monkeypatch.setattr(mfa_gate, "_request_json", fake_request)
        mfa_gate._sync_casbin_after_seed(opener, cookie_jar, "http://backend.local")
        assert calls[0][0][1] == "POST"
        assert calls[0][0][2].endswith(mfa_gate.CASBIN_SYNC_ROUTE)
        assert calls[0][1]["extra_headers"] == {"X-CSRF-Token": "csrf-value"}

        monkeypatch.setattr(
            mfa_gate,
            "_request_json",
            lambda *args, **kwargs: (
                200,
                {
                    "dry_run": False,
                    "results": _ket_qua_sync_day_du(
                        mfa_gate,
                        **{"role:officer": {"success": False, "error": "boom"}},
                    ),
                },
            ),
        )
        with pytest.raises(mfa_gate.GateError, match="role:officer"):
            mfa_gate._sync_casbin_after_seed(opener, cookie_jar, "http://backend.local")

    # ------------------------------------------------------------------
    # Kết quả sync phải TỰ CHỨNG MINH nó hội tụ.
    #
    # Phép cũ chỉ hỏi "có vai nào fail không". Một map `results` RỖNG không có
    # vai nào để fail ⇒ xanh; thiếu vai ⇒ xanh; `{"success": None}` ⇒ xanh.
    # Mỗi ca dưới đây vi phạm ĐÚNG MỘT thứ, xuất phát từ một map hợp lệ.
    # ------------------------------------------------------------------

    @staticmethod
    def _sync_voi(mfa_gate, monkeypatch, results):
        """Chạy `_sync_casbin_after_seed` với một map `results` cho trước."""
        jar = mfa_gate.http.cookiejar.CookieJar()
        jar.set_cookie(
            mfa_gate.http.cookiejar.Cookie(
                version=0, name="csrf_token", value="csrf-value", port=None,
                port_specified=False, domain="backend.local", domain_specified=True,
                domain_initial_dot=False, path="/", path_specified=True, secure=False,
                expires=None, discard=True, comment=None, comment_url=None, rest={},
                rfc2109=False,
            )
        )
        monkeypatch.setattr(
            mfa_gate,
            "_request_json",
            lambda *args, **kwargs: (200, {"dry_run": False, "results": results}),
        )
        mfa_gate._sync_casbin_after_seed(
            mfa_gate.urllib.request.build_opener(), jar, "http://backend.local"
        )

    def test_sync_casbin_khop_system_roles_cua_san_pham(self, mfa_gate):
        """`EXPECTED_SYNC_ROLES` phải BẰNG `SYSTEM_ROLES` của sản phẩm.

        Cổng đòi đủ sáu vai; nếu danh sách ấy chép tay và sản phẩm thêm vai
        thứ bảy thì lượt sync bỏ sót đúng vai mới mà cổng vẫn xanh.
        """
        assert set(mfa_gate.EXPECTED_SYNC_ROLES) == set(_system_roles_san_pham()), (
            "EXPECTED_SYNC_ROLES lệch SYSTEM_ROLES: cổng=%s sản phẩm=%s"
            % (sorted(mfa_gate.EXPECTED_SYNC_ROLES), sorted(_system_roles_san_pham()))
        )
        assert len(set(mfa_gate.EXPECTED_SYNC_ROLES)) == len(
            mfa_gate.EXPECTED_SYNC_ROLES
        ), "EXPECTED_SYNC_ROLES có tên trùng"

    def test_sync_casbin_ban_hop_le_phai_xanh(self, mfa_gate, monkeypatch):
        """Đối chứng XANH: đủ sáu vai, admin skip-vì-an-toàn, còn lại success.

        Không có ca này thì mọi ca đỏ bên dưới có thể đang đỏ vì cổng đã siết
        quá tay, chứ không phải vì bắt đúng thứ nó nhắm.
        """
        self._sync_voi(mfa_gate, monkeypatch, _ket_qua_sync_day_du(mfa_gate))

        # Bỏ qua vì "không lệch" cũng là kết cục hợp lệ cho vai không phải admin.
        self._sync_voi(
            mfa_gate,
            monkeypatch,
            _ket_qua_sync_day_du(
                mfa_gate,
                **{"role:user": {"skipped": True, "reason": "No drift detected"}},
            ),
        )

    @pytest.mark.parametrize(
        "mo_ta,results_fn,mau",
        [
            (
                "results rỗng",
                lambda g: {},
                "EMPTY results map",
            ),
            (
                "thiếu một vai",
                lambda g: {
                    k: v
                    for k, v in _ket_qua_sync_day_du(g).items()
                    if k != "role:collaborator"
                },
                "role:collaborator",
            ),
            (
                "thừa một vai lạ",
                lambda g: _ket_qua_sync_day_du(
                    g, **{"role:ghost": {"success": True}}
                ),
                "role:ghost",
            ),
            (
                "vai trả dict RỖNG",
                lambda g: _ket_qua_sync_day_du(g, **{"role:officer": {}}),
                "empty/invalid result object",
            ),
            (
                "success=False",
                lambda g: _ket_qua_sync_day_du(
                    g, **{"role:officer": {"success": False}}
                ),
                "ambiguous",
            ),
            (
                "success=None",
                lambda g: _ket_qua_sync_day_du(
                    g, **{"role:officer": {"success": None}}
                ),
                "ambiguous",
            ),
            (
                "success=0",
                lambda g: _ket_qua_sync_day_du(g, **{"role:officer": {"success": 0}}),
                "ambiguous",
            ),
            (
                'success="false" (chuỗi truthy)',
                lambda g: _ket_qua_sync_day_du(
                    g, **{"role:officer": {"success": "false"}}
                ),
                "ambiguous",
            ),
            (
                "skipped=False",
                lambda g: _ket_qua_sync_day_du(
                    g, **{"role:officer": {"skipped": False, "reason": "x"}}
                ),
                "ambiguous",
            ),
            (
                "vừa success vừa skipped",
                lambda g: _ket_qua_sync_day_du(
                    g,
                    **{
                        "role:officer": {
                            "success": True,
                            "skipped": True,
                            "reason": "No drift detected",
                        }
                    },
                ),
                "ambiguous",
            ),
            (
                "success=True kèm error khác rỗng",
                lambda g: _ket_qua_sync_day_du(
                    g, **{"role:officer": {"success": True, "error": "boom"}}
                ),
                "reported an error",
            ),
            (
                "skipped nhưng KHÔNG có reason",
                lambda g: _ket_qua_sync_day_du(
                    g, **{"role:officer": {"skipped": True}}
                ),
                "without a reason",
            ),
            (
                "skipped với reason rỗng/trắng",
                lambda g: _ket_qua_sync_day_du(
                    g, **{"role:officer": {"skipped": True, "reason": "   "}}
                ),
                "without a reason",
            ),
            (
                "admin báo success thay vì skip-vì-an-toàn",
                lambda g: _ket_qua_sync_day_du(
                    g, **{"role:admin": {"success": True}}
                ),
                "must stay skipped",
            ),
            (
                "admin skip nhưng SAI lý do",
                lambda g: _ket_qua_sync_day_du(
                    g,
                    **{
                        "role:admin": {
                            "skipped": True,
                            "reason": "No drift detected",
                        }
                    },
                ),
                "wrong reason",
            ),
            (
                'vai thường trả "No template defined"',
                lambda g: _ket_qua_sync_day_du(
                    g,
                    **{
                        "role:officer": {
                            "skipped": True,
                            "reason": "No template defined",
                        }
                    },
                ),
                "unacceptable reason",
            ),
        ],
    )
    def test_sync_casbin_ket_qua_fail_closed(
        self, mfa_gate, monkeypatch, mo_ta, results_fn, mau
    ):
        """Mười sáu đường mà lượt sync KHÔNG hội tụ đều phải ĐỎ."""
        with pytest.raises(mfa_gate.GateError, match=re.escape(mau)) as loi:
            self._sync_voi(mfa_gate, monkeypatch, results_fn(mfa_gate))
        assert str(loi.value), mo_ta

    def test_sync_casbin_thieu_csrf_phai_do_truoc_request(
        self, mfa_gate, monkeypatch
    ):
        monkeypatch.setattr(
            mfa_gate,
            "_request_json",
            lambda *args, **kwargs: pytest.fail(
                "không được gọi request khi thiếu CSRF"
            ),
        )
        with pytest.raises(mfa_gate.GateError, match="did not issue the CSRF cookie"):
            mfa_gate._sync_casbin_after_seed(
                mfa_gate.urllib.request.build_opener(),
                mfa_gate.http.cookiejar.CookieJar(),
                "http://backend.local",
            )

    def test_login_200_khong_challenge_khong_duoc_tinh_la_mfa(
        self, mfa_gate, monkeypatch
    ):
        monkeypatch.setattr(
            mfa_gate,
            "_request_json",
            lambda *args, **kwargs: (200, {"user": {"role": "admin"}}),
        )
        with pytest.raises(
            mfa_gate.GateError, match="without the required MFA challenge"
        ):
            mfa_gate._preflight_account(
                (
                    "admin",
                    "admin",
                    "password",
                    "admin",
                    "JBSWY3DPEHPK3PXPJBSWY3DPEHPK3PXP",
                ),
                "http://localhost:8000",
            )

    def test_login_va_verify_200_nhung_route_casbin_403_phai_do(
        self, mfa_gate, monkeypatch
    ):
        responses = iter(
            [
                (200, {"mfa_required": True, "mfa_token": "challenge"}),
                (401, {"detail": "Not authenticated"}),  # probe TRƯỚC verify
                (200, {"user": {"role": "manager"}}),
                (403, {"detail": "forbidden"}),
            ]
        )
        monkeypatch.setattr(
            mfa_gate, "_request_json", lambda *args, **kwargs: next(responses)
        )
        monkeypatch.setattr(
            mfa_gate, "_totp_for_preflight", lambda *args: ("123456", 55_000_000)
        )
        with pytest.raises(mfa_gate.GateError, match=r"pipeline/all \(HTTP 403\)"):
            mfa_gate._preflight_account(
                (
                    "manager",
                    "manager",
                    "password",
                    "manager",
                    "JBSWY3DPEHPK3PXPJBSWY3DPEHPK3PXP",
                ),
                "http://localhost:8000",
            )

    def test_verify_mfa_401_phai_do_va_khong_duoc_probe_casbin(
        self, mfa_gate, monkeypatch
    ):
        calls = []
        responses = iter(
            [
                (200, {"mfa_required": True, "mfa_token": "challenge"}),
                (401, {"detail": "Not authenticated"}),  # probe TRƯỚC verify
                (401, {"detail": "Invalid verification code"}),
            ]
        )

        def fake_request(*args, **kwargs):
            calls.append((args, kwargs))
            return next(responses)

        monkeypatch.setattr(mfa_gate, "_request_json", fake_request)
        monkeypatch.setattr(
            mfa_gate, "_totp_for_preflight", lambda *args: ("123456", 55_000_000)
        )
        with pytest.raises(mfa_gate.GateError, match="verification returned HTTP 401"):
            mfa_gate._preflight_account(
                (
                    "admin",
                    "admin",
                    "password",
                    "admin",
                    "JBSWY3DPEHPK3PXPJBSWY3DPEHPK3PXP",
                ),
                "http://localhost:8000",
            )
        # Ba lượt: login → probe TRƯỚC verify (phải 401) → verify (đỏ).
        # Điều phải chứng minh là KHÔNG có lượt thứ tư: probe Casbin SAU verify.
        # So `len(calls)` trần thì mơ hồ — nó cũng xanh nếu thứ tự bị đảo.
        buoc = [(a[1], a[2]) for a, _ in calls]
        assert [m for m, _ in buoc] == ["POST", "GET", "POST"], (
            "trình tự sai: %r" % (buoc,)
        )
        assert buoc[1][1].endswith(mfa_gate.PROTECTED_ROUTE)
        assert buoc[2][1].endswith("/api/auth/verify-mfa")
        sau_verify = buoc[3:]
        assert not sau_verify, (
            "verify MFA đỏ mà cổng vẫn gọi tiếp: %r" % (sau_verify,)
        )

    def test_officer_bi_bat_mfa_phai_do(self, mfa_gate, monkeypatch):
        monkeypatch.setattr(
            mfa_gate,
            "_request_json",
            lambda *args, **kwargs: (200, {"mfa_required": True, "mfa_token": "x"}),
        )
        with pytest.raises(mfa_gate.GateError, match="unexpectedly requires MFA"):
            mfa_gate._preflight_account(
                ("officer", "officer", "password", "officer", None),
                "http://localhost:8000",
            )

    # ------------------------------------------------------------------
    # Chặng chỉ-mật-khẩu KHÔNG được là một phiên.
    #
    # `/api/auth/login` trả HTTP 200 cho CẢ hai nhánh — đăng nhập xong và mới
    # qua yếu tố thứ nhất. Một cổng chỉ kiểm "200" vì thế XANH y hệt khi MFA bị
    # vô hiệu hoá. Bốn ca dưới đây canh bốn đường rò khác nhau của cùng một
    # bất biến, mỗi ca vi phạm ĐÚNG MỘT thứ.
    # ------------------------------------------------------------------

    def _chay_preflight_admin(self, mfa_gate, monkeypatch, responses, jar=None):
        """Chạy `_preflight_account` cho admin với chuỗi phản hồi cho trước."""
        calls = []

        def fake_request(opener, method, url, **kwargs):
            calls.append((method, url, kwargs))
            return responses.pop(0)

        monkeypatch.setattr(mfa_gate, "_request_json", fake_request)
        monkeypatch.setattr(
            mfa_gate, "_totp_for_preflight", lambda *args: ("123456", 55_000_000)
        )
        if jar is not None:
            monkeypatch.setattr(
                mfa_gate.http.cookiejar, "CookieJar", lambda *a, **k: jar
            )
        mfa_gate._preflight_account(
            ("admin", "admin", "password", "admin", "JBSWY3DPEHPK3PXPJBSWY3DPEHPK3PXP"),
            "http://localhost:8000",
        )
        return calls

    def test_pre_verify_route_200_la_ro_phien_phai_do(self, mfa_gate, monkeypatch):
        """Route Casbin trả 200 TRƯỚC verify ⇒ phiên đã được cấp bởi mật khẩu."""
        with pytest.raises(mfa_gate.GateError) as loi:
            self._chay_preflight_admin(
                mfa_gate,
                monkeypatch,
                [
                    (200, {"mfa_required": True, "mfa_token": "challenge"}),
                    (200, []),  # <-- lẽ ra phải 401
                ],
            )
        assert "BEFORE MFA verification" in str(loi.value)
        assert "expected exactly 401" in str(loi.value)

    def test_pre_verify_route_403_cung_phai_do(self, mfa_gate, monkeypatch):
        """403 KHÔNG được coi là đạt.

        403 nghĩa là máy chủ đã nhận ra một danh tính rồi mới chặn quyền — tức
        phiên ĐÃ tồn tại. Chỉ 401 (chưa xác thực) mới chứng minh chặng
        chỉ-mật-khẩu không cấp gì.
        """
        with pytest.raises(mfa_gate.GateError) as loi:
            self._chay_preflight_admin(
                mfa_gate,
                monkeypatch,
                [
                    (200, {"mfa_required": True, "mfa_token": "challenge"}),
                    (403, {"detail": "MFA is required for privileged accounts."}),
                ],
            )
        assert "expected exactly 401" in str(loi.value)

    @pytest.mark.parametrize(
        "gia_tri",
        ["leaked-value", "", None, 0],
        ids=["chuoi-that", "chuoi-rong", "None", "so-khong"],
    )
    def test_challenge_ro_token_trong_than_phai_do(
        self, mfa_gate, monkeypatch, gia_tri
    ):
        """Thân challenge mang access/refresh token ⇒ rò phiên.

        Bốn giá trị, không phải một: `""`, `None`, `0` đều FALSY, nên một cổng
        kiểm `login.get(khoa)` sẽ cho chúng đi lọt. Chính việc khoá CÓ MẶT đã
        là bằng chứng chặng chỉ-mật-khẩu đang cấp phát trường phiên — giá trị
        của nó không đổi được điều đó.
        """
        for khoa in ("access_token", "refresh_token"):
            with pytest.raises(mfa_gate.GateError) as loi:
                self._chay_preflight_admin(
                    mfa_gate,
                    monkeypatch,
                    [
                        (
                            200,
                            {
                                "mfa_required": True,
                                "mfa_token": "challenge",
                                khoa: gia_tri,
                            },
                        ),
                    ],
                )
            assert khoa in str(loi.value)
            assert "before the second factor" in str(loi.value)

    def test_challenge_dat_cookie_phien_phai_do(self, mfa_gate, monkeypatch):
        """Challenge đặt cookie phiên ⇒ rò, dù thân phản hồi sạch."""

        class _CookieGia:
            def __init__(self, name):
                self.name = name

        jar = [_CookieGia("access_token")]
        with pytest.raises(mfa_gate.GateError) as loi:
            self._chay_preflight_admin(
                mfa_gate,
                monkeypatch,
                [(200, {"mfa_required": True, "mfa_token": "challenge"})],
                jar=jar,
            )
        assert "session cookie" in str(loi.value)
        assert "access_token" in str(loi.value)

    def test_cookie_csrf_khong_bi_tinh_la_ro_phien(self, mfa_gate, monkeypatch):
        """Đối chứng: `csrf_token` KHÔNG phải bằng chứng xác thực.

        Nếu ca này đỏ nghĩa là `COOKIE_PHIEN` bị nới quá tay và cổng sẽ báo
        động giả trên một cookie hoàn toàn hợp lệ.
        """
        assert "csrf_token" not in mfa_gate.COOKIE_PHIEN

        class _CookieGia:
            def __init__(self, name):
                self.name = name

        calls = self._chay_preflight_admin(
            mfa_gate,
            monkeypatch,
            [
                (200, {"mfa_required": True, "mfa_token": "challenge"}),
                (401, {"detail": "Not authenticated"}),
                (200, {"user": {"role": "admin"}}),
                (200, []),
            ],
            jar=[_CookieGia("csrf_token")],
        )
        assert [c[0] for c in calls] == ["POST", "GET", "POST", "GET"]

    # ------------------------------------------------------------------
    # Đọc lại sau commit phải đi qua PHIÊN MỚI.
    # ------------------------------------------------------------------

    def test_bootstrap_doc_lai_bang_session_moi_sau_commit(self, mfa_gate):
        """`expire_on_commit=False` làm phép "đọc lại" trong cùng session vô nghĩa.

        `app/database.py` khai `AsyncSessionLocal(..., expire_on_commit=False)`,
        nên sau `commit()` các đối tượng vẫn nằm trong identity map với giá trị
        TRONG BỘ NHỚ. Một `select()` chạy lại trên chính session ấy trả về đúng
        những instance đó ⇒ phép kiểm persistence XANH kể cả khi không dòng nào
        chạm đĩa. Chỉ một phiên mới (identity map rỗng) mới buộc phát SELECT thật.
        """
        ghi, kiem = _hai_phien_bootstrap()

        # (a) Hai context KHÁC NHAU, alias khác nhau.
        assert ghi.alias != kiem.alias, (
            "chỉ có MỘT phiên: đọc lại sẽ trúng identity map của chính nó"
        )
        assert kiem.alias == "kiem", "phiên kiểm phải mang alias `kiem`, thấy %r" % (
            kiem.alias,
        )

        # (b) KHÔNG lồng nhau — phiên kiểm nằm SAU khi phiên ghi kết thúc.
        assert kiem.node.lineno > ghi.ket_thuc, (
            "phiên kiểm (dòng %d) nằm BÊN TRONG phiên ghi (kết thúc dòng %d) — "
            "phiên ghi chưa đóng thì đọc lại vẫn có thể trúng bộ nhớ của nó."
            % (kiem.node.lineno, ghi.ket_thuc)
        )

        # (c) Commit phải là ĐÚNG `await db.commit()`.
        #
        #     Phép cũ chỉ tìm một `ast.Attribute` tên `commit`, nên ba đột biến
        #     đi lọt trong khi KHÔNG có gì được ghi xuống đĩa:
        #       - `db.commit`                   — thiếu dấu gọi, chỉ là biểu thức
        #       - `await officer_user.commit()` — commit nhầm chủ thể
        #       - `await db.commit(<gì đó>)`    — không phải chữ ký thật
        cau_commit = [n for n in ast.walk(ghi.node) if _la_await_commit(n, ghi.alias)]
        assert len(cau_commit) == 1, (
            "cần ĐÚNG một câu lệnh `await %s.commit()` trong phiên ghi, thấy %d"
            % (ghi.alias, len(cau_commit))
        )
        dong_commit = cau_commit[0].lineno

        # Mọi tham chiếu `.commit` KHÁC trong phiên ghi đều bị loại: nó khiến
        # người đọc tin là đã commit trong khi đường chạy thật thì không.
        hop_le = {n for n in ast.walk(cau_commit[0]) if isinstance(n, ast.Attribute)}
        commit_la = sorted(
            n.lineno
            for n in ast.walk(ghi.node)
            if isinstance(n, ast.Attribute) and n.attr == "commit" and n not in hop_le
        )
        assert not commit_la, (
            "phiên ghi còn tham chiếu `.commit` khác ở dòng %s — chỉ được đúng một "
            "`await %s.commit()`" % (commit_la, ghi.alias)
        )

        # Commit phải đứng SAU mọi lời gọi `enable_mfa`. Commit trước thì đúng
        # những hàng vừa bật MFA lại nằm ngoài transaction đã kết.
        dong_enable = sorted(
            n.lineno
            for n in ast.walk(ghi.node)
            if isinstance(n, ast.Call)
            and isinstance(n.func, ast.Attribute)
            and n.func.attr == "enable_mfa"
        )
        assert dong_enable, "phiên ghi không gọi `enable_mfa` — nó không bật gì cả"
        assert max(dong_enable) < dong_commit, (
            "`await %s.commit()` (dòng %d) đứng TRƯỚC `enable_mfa` (dòng %s)"
            % (ghi.alias, dong_commit, dong_enable)
        )
        assert dong_commit < kiem.node.lineno, (
            "commit (dòng %d) không đứng trước phiên kiểm (dòng %d)"
            % (dong_commit, kiem.node.lineno)
        )

        # (d) MỌI `.execute()` trong phiên kiểm phải gọi trên `kiem`.
        goi = _alias_cua_execute(kiem.node)
        assert goi, "phiên kiểm không phát `.execute()` nào — nó không đọc gì cả"
        sai = sorted(a for a in goi if a != "kiem")
        assert not sai, (
            "phiên kiểm gọi `.execute()` trên %s thay vì `kiem` — mở phiên mới mà "
            "vẫn truy vấn bằng phiên cũ thì identity map cũ vẫn được dùng." % sai
        )

        # (e) Không còn `close()` tay trên phiên ghi: `rollback()` ở nhánh lỗi sẽ
        #     áp lên một session đã đóng.
        dong_close = [
            n
            for n in ast.walk(ghi.node)
            if isinstance(n, ast.Attribute)
            and n.attr == "close"
            and isinstance(n.value, ast.Name)
            and n.value.id == ghi.alias
        ]
        assert not dong_close, (
            "phiên ghi bị `close()` bằng tay bên trong `async with` — vòng đời bẩn"
        )

    def test_bootstrap_kiem_du_role_status_va_ciphertext(self, mfa_gate):
        """Phiên mới phải kiểm đủ bốn thứ, không chỉ cờ `mfa_enabled`."""
        nguon = _phai_ton_tai(DUONG_MFA_GATE, "script cổng MFA nightly").read_text(
            encoding="utf-8"
        )
        _, kiem = _hai_phien_bootstrap()
        # Cắt theo PHẠM VI AST của phiên kiểm, không cắt theo chuỗi: cắt chuỗi
        # sẽ nuốt luôn mọi thứ phía sau hàm và làm phép kiểm xanh nhờ mã ở nơi khác.
        khoi = "\n".join(
            nguon.split("\n")[kiem.node.lineno - 1 : kiem.ket_thuc]
        )
        for can, vi_sao in (
            ('user.status != "active"', "trạng thái tài khoản"),
            ("user.role != expected_role", "vai trò"),
            ("user.mfa_enabled is not True", "cờ MFA"),
            ("user.totp_secret_encrypted", "ciphertext tồn tại"),
            ("decrypt_secret", "giải mã round-trip"),
            ("compare_digest", "so sánh secret theo thời gian hằng"),
            ("officer_after.mfa_enabled", "officer KHÔNG được bật MFA"),
            ("enabled_names != set(expected_privileged)", "tập MFA-enabled CHÍNH XÁC"),
        ):
            assert can in khoi, "khối kiểm phiên mới thiếu phép %s (%r)" % (vi_sao, can)

    # ------------------------------------------------------------------
    # Hai lỗ nối dây.
    # ------------------------------------------------------------------

    def test_sync_casbin_la_buoc_rieng_dung_danh_tinh_admin(
        self, mfa_gate, monkeypatch, tmp_path
    ):
        """`sync-casbin` phải: đúng admin, đúng phiên đã xác thực, ĐÚNG một lượt.

        Trước đây lượt sync là một callback móc vào preflight. Nó phải là bước
        riêng vì endpoint sync chỉ nạp lại enforcer của MỘT worker — nó là bước
        GHI CSDL, và bước cutover sau nó mới đưa cả fleet về cùng ảnh chụp.
        """
        tep = tmp_path / "counter"
        monkeypatch.setenv("E2E_API_URL", "http://localhost:8000/")
        monkeypatch.setenv("QLTS_TOTP_COUNTER_FILE", str(tep))
        monkeypatch.setattr(
            mfa_gate,
            "_runtime_account_values",
            lambda: [
                ("admin", "admin", "pw", "admin", "S1"),
                ("officer", "vothithuthuhien", "pw", "officer", None),
                ("manager", "phanthithuyvan", "pw", "manager", "S2"),
                ("setup", "vothithuthuhien", "pw", "officer", None),
            ],
        )

        da_dang_nhap: list[tuple] = []
        phien_gia = mfa_gate.PhienDaXacThuc(
            label="admin",
            opener=object(),
            cookie_jar=object(),
            than={"user": {"role": "admin"}},
            counter_da_tieu=55_000_123,
        )

        def fake_dang_nhap(account, base_url):
            da_dang_nhap.append((account, base_url))
            return phien_gia

        sync: list[tuple] = []
        monkeypatch.setattr(mfa_gate, "_dang_nhap_va_chung_minh", fake_dang_nhap)
        monkeypatch.setattr(
            mfa_gate,
            "_sync_casbin_after_seed",
            lambda opener, jar, base: sync.append((opener, jar, base)),
        )

        mfa_gate.sync_casbin()

        # ĐÚNG một lượt đăng nhập, và là admin — không phải tài khoản nào khác.
        assert len(da_dang_nhap) == 1, "sync đăng nhập %d lần" % len(da_dang_nhap)
        assert da_dang_nhap[0][0][0] == "admin"
        assert da_dang_nhap[0][1] == "http://localhost:8000", "base_url chưa rstrip"

        # Sync chạy đúng một lần, với CHÍNH opener/jar của phiên vừa xác thực.
        assert len(sync) == 1, "sync Casbin chạy %d lần, cần đúng 1" % len(sync)
        assert sync[0][0] is phien_gia.opener
        assert sync[0][1] is phien_gia.cookie_jar
        assert sync[0][2] == "http://localhost:8000"

        # Counter đã đốt phải được ghi ra tệp — chỉ MỘT số, không gì khác.
        assert tep.read_text(encoding="utf-8").strip() == "55000123"

    def test_sync_casbin_ghi_counter_TRUOC_khi_goi_sync(
        self, mfa_gate, monkeypatch, tmp_path
    ):
        """Sync đỏ thì counter VẪN đã bị đốt thật ⇒ vẫn phải để lại dấu.

        Ghi counter sau khi sync thành công là một lỗ: lượt sync hỏng vẫn tiêu
        một counter của backend, mà bước sau lại không biết để chờ qua nó.
        """
        tep = tmp_path / "counter"
        monkeypatch.setenv("E2E_API_URL", "http://localhost:8000")
        monkeypatch.setenv("QLTS_TOTP_COUNTER_FILE", str(tep))
        monkeypatch.setattr(
            mfa_gate,
            "_runtime_account_values",
            lambda: [("admin", "admin", "pw", "admin", "S1")],
        )
        monkeypatch.setattr(
            mfa_gate,
            "_dang_nhap_va_chung_minh",
            lambda a, b: mfa_gate.PhienDaXacThuc(
                "admin", object(), object(), {}, 55_000_456
            ),
        )

        def sync_hong(opener, jar, base):
            raise mfa_gate.GateError("sync hỏng")

        monkeypatch.setattr(mfa_gate, "_sync_casbin_after_seed", sync_hong)
        with pytest.raises(mfa_gate.GateError, match="sync hỏng"):
            mfa_gate.sync_casbin()
        assert tep.read_text(encoding="utf-8").strip() == "55000456"

    def test_sync_casbin_doi_tai_khoan_co_MFA(self, mfa_gate, monkeypatch, tmp_path):
        """Tài khoản sync không đi qua MFA ⇒ không có counter nào để cách ly."""
        monkeypatch.setenv("E2E_API_URL", "http://localhost:8000")
        monkeypatch.setenv("QLTS_TOTP_COUNTER_FILE", str(tmp_path / "c"))
        monkeypatch.setattr(
            mfa_gate,
            "_runtime_account_values",
            lambda: [("admin", "admin", "pw", "admin", None)],
        )
        monkeypatch.setattr(
            mfa_gate,
            "_dang_nhap_va_chung_minh",
            lambda a, b: mfa_gate.PhienDaXacThuc("admin", object(), object(), {}, None),
        )
        monkeypatch.setattr(
            mfa_gate,
            "_sync_casbin_after_seed",
            lambda *a: pytest.fail("không được sync"),
        )
        with pytest.raises(mfa_gate.GateError, match="không đi qua MFA"):
            mfa_gate.sync_casbin()

    # ------------------------------------------------------------------
    # Cách ly counter TOTP giữa sync và preflight.
    # ------------------------------------------------------------------

    def test_tep_counter_chi_chua_so(self, mfa_gate, tmp_path):
        """Tệp trung gian chỉ được mang MỘT số nguyên, không gì khác.

        Nó đi qua ranh giới hai bước workflow nên là một kênh dữ liệu thật.
        Mọi thứ khác lọt vào đây — secret, mã TOTP, tên tài khoản — là bản sao
        bí mật nằm ngoài mọi cơ chế che.
        """
        tep = tmp_path / "counter"
        mfa_gate._ghi_counter_da_tieu(tep, 55_000_000)
        assert tep.read_text(encoding="utf-8") == "55000000\n"
        assert mfa_gate._doc_counter_da_tieu(tep) == 55_000_000

        for xau in ("", "  ", "abc", "55000000 JBSWY3DP", "-1", "5.5", "0x10"):
            tep.write_text(xau, encoding="utf-8")
            with pytest.raises(mfa_gate.GateError, match="chỉ chứa chữ số"):
                mfa_gate._doc_counter_da_tieu(tep)

    def test_thieu_tep_counter_phai_do_khong_duoc_mac_dinh_0(self, mfa_gate, tmp_path):
        """Thiếu tệp ⇒ ĐỎ. Mặc định 0 sẽ làm phép chờ luôn thoả ngay lập tức."""
        with pytest.raises(mfa_gate.GateError, match="thiếu tệp counter"):
            mfa_gate._doc_counter_da_tieu(tmp_path / "khong-ton-tai")

    def test_ghi_counter_tu_choi_gia_tri_khong_hop_le(self, mfa_gate, tmp_path):
        """`True` là `int` trong Python — phải bị loại, nếu không ghi ra `1`."""
        tep = tmp_path / "c"
        for xau in (True, -1, "55", 5.0, None):
            with pytest.raises(mfa_gate.GateError, match="số nguyên không âm"):
                mfa_gate._ghi_counter_da_tieu(tep, xau)

    def test_dieu_kien_cho_counter_phai_nghiem_ngat(self, mfa_gate, monkeypatch):
        """`>` chứ không `>=`: backend từ chối khi `counter <= đã_lưu`.

        Ca biên là lúc `counter_hiện_tại - 1` BẰNG counter đã tiêu. Với `>=`
        phép chờ trả về ngay, preflight gửi lại đúng mã đã đốt, và verify-mfa
        bị guard đơn điệu từ chối — trong khi thiết kế cấm retry.
        """
        da_tieu = 55_000_000
        # Đồng hồ đứng yên ở đúng ca biên: floor(t/30) - 1 == da_tieu.
        monkeypatch.setattr(time, "time", lambda: (da_tieu + 1) * 30 + 5)
        monkeypatch.setattr(mfa_gate.time, "sleep", lambda _s: None)
        with pytest.raises(mfa_gate.GateError, match="chưa vượt"):
            mfa_gate._cho_counter_vuot(da_tieu, han_giay=0.0)

        # Vượt đúng một bậc thì phải trả về ngay.
        monkeypatch.setattr(time, "time", lambda: (da_tieu + 2) * 30 + 5)
        assert mfa_gate._cho_counter_vuot(da_tieu, han_giay=0.0) == da_tieu + 2

    def test_cho_counter_het_han_phai_do_khong_duoc_di_tiep(
        self, mfa_gate, monkeypatch
    ):
        """Hết hạn mà đồng hồ không tiến là sự cố thật — ĐỎ, không thử lại."""
        monkeypatch.setattr(time, "time", lambda: 55_000_000 * 30 + 1)
        monkeypatch.setattr(mfa_gate.time, "sleep", lambda _s: None)
        with pytest.raises(mfa_gate.GateError, match="đồng hồ không tiến"):
            mfa_gate._cho_counter_vuot(55_000_000, han_giay=0.0)

    # ------------------------------------------------------------------
    # Preflight cuối: CHỈ ĐỌC, và lấy mẫu đủ số kết nối.
    # ------------------------------------------------------------------

    def _preflight_gia_lap(self, mfa_gate, monkeypatch, tmp_path, *, ma_probe=200):
        """Chạy `preflight()` thật với mạng giả lập; trả danh sách URL đã gọi."""
        tep = tmp_path / "counter"
        tep.write_text("10\n", encoding="utf-8")
        monkeypatch.setenv("E2E_API_URL", "http://localhost:8000")
        monkeypatch.setenv("QLTS_TOTP_COUNTER_FILE", str(tep))
        monkeypatch.setattr(
            mfa_gate,
            "_runtime_account_values",
            lambda: [
                ("admin", "admin", "pw", "admin", "S1"),
                ("officer", "vothithuthuhien", "pw", "officer", None),
                ("manager", "phanthithuyvan", "pw", "manager", "S2"),
                ("setup", "vothithuthuhien", "pw", "officer", None),
            ],
        )
        monkeypatch.setattr(mfa_gate, "_cho_counter_vuot", lambda *a, **k: 12)
        monkeypatch.setattr(
            mfa_gate,
            "_dang_nhap_va_chung_minh",
            lambda acc, base: mfa_gate.PhienDaXacThuc(
                acc[0], object(), object(), {"user": {"role": acc[3]}}, None
            ),
        )
        goi: list[tuple] = []

        def fake_request(opener, method, url, **kwargs):
            goi.append((method, url))
            return (ma_probe, [])

        monkeypatch.setattr(mfa_gate, "_request_json", fake_request)
        return goi

    def test_preflight_cho_counter_truoc_moi_cham_mang(
        self, mfa_gate, monkeypatch, tmp_path
    ):
        """`preflight()` phải THỰC SỰ gọi phép chờ, và gọi TRƯỚC mọi request.

        Ca này tồn tại vì đột biến "gỡ hẳn lời gọi `_cho_counter_vuot`" đi lọt
        qua toàn bộ bộ test: mọi ca khác đều monkeypatch phép chờ, nên chúng
        xanh y hệt dù đường chạy thật không còn chờ nữa. Không chờ thì mã của
        `counter trước` có thể vẫn đúng counter mà bước sync đã đốt, và backend
        từ chối nó theo bất biến đơn điệu — trong khi thiết kế CẤM retry.
        """
        tep = tmp_path / "counter"
        tep.write_text("777\n", encoding="utf-8")
        monkeypatch.setenv("E2E_API_URL", "http://localhost:8000")
        monkeypatch.setenv("QLTS_TOTP_COUNTER_FILE", str(tep))
        monkeypatch.setattr(
            mfa_gate,
            "_runtime_account_values",
            lambda: [("officer", "vothithuthuhien", "pw", "officer", None)],
        )

        su_kien: list[str] = []
        monkeypatch.setattr(
            mfa_gate,
            "_cho_counter_vuot",
            lambda counter, **k: su_kien.append("CHO:%d" % counter) or 999,
        )
        monkeypatch.setattr(
            mfa_gate,
            "_dang_nhap_va_chung_minh",
            lambda acc, base: su_kien.append("DANGNHAP")
            or mfa_gate.PhienDaXacThuc(acc[0], object(), object(), {}, None),
        )
        monkeypatch.setattr(
            mfa_gate,
            "_request_json",
            lambda *a, **k: su_kien.append("HTTP") or (200, []),
        )
        monkeypatch.setattr(mfa_gate, "_probe_fleet", lambda *a, **k: None)

        with pytest.raises(mfa_gate.GateError, match="thiếu phiên"):
            mfa_gate.preflight()

        assert "CHO:777" in su_kien, "preflight KHÔNG chờ counter đã tiêu"
        assert su_kien[0] == "CHO:777", (
            "phép chờ phải đứng TRƯỚC mọi lượt chạm mạng, thấy %r" % (su_kien[:3],)
        )

    def test_preflight_read_only_khong_goi_route_ghi(
        self, mfa_gate, monkeypatch, tmp_path
    ):
        """Preflight cuối không được chạm endpoint sync — nó chỉ ĐỌC.

        Móc sync vào đây nghĩa là cái 200 đo được một phần do CHÍNH nó vừa tạo
        ra: cổng tự chứng minh mình bằng thứ mình vừa thay đổi.
        """
        goi = self._preflight_gia_lap(mfa_gate, monkeypatch, tmp_path)
        monkeypatch.setattr(
            mfa_gate,
            "_sync_casbin_after_seed",
            lambda *a, **k: pytest.fail("preflight KHÔNG được sync policy"),
        )
        mfa_gate.preflight()

        assert all(m == "GET" for m, _u in goi), "preflight phát request không phải GET"
        assert all(mfa_gate.CASBIN_SYNC_ROUTE.split("?")[0] not in u for _m, u in goi)
        assert all(u.endswith(mfa_gate.PROTECTED_ROUTE) for _m, u in goi)

    def test_preflight_lay_mau_du_luot_cho_ba_vai_dac_quyen(
        self, mfa_gate, monkeypatch, tmp_path
    ):
        """4 probe của 4 đường + 3 × PROBE_FLEET_LUOT kết nối lấy mẫu."""
        goi = self._preflight_gia_lap(mfa_gate, monkeypatch, tmp_path)
        mfa_gate.preflight()
        assert len(goi) == 4 + 3 * mfa_gate.PROBE_FLEET_LUOT, len(goi)
        assert mfa_gate.PROBE_FLEET_LUOT >= 32

    def test_probe_fleet_mot_ma_khac_200_la_do(self, mfa_gate, monkeypatch):
        """Chỉ MỘT lượt 403 giữa 32 lượt cũng phải ĐỎ — fleet chưa hội tụ."""
        phien = mfa_gate.PhienDaXacThuc("manager", object(), object(), {}, None)
        con_lai = [200] * (mfa_gate.PROBE_FLEET_LUOT - 1)
        thu_tu = con_lai[:5] + [403] + con_lai[5:]

        def fake_request(opener, method, url, **kwargs):
            return (thu_tu.pop(0), {"detail": "forbidden"})

        monkeypatch.setattr(mfa_gate, "_request_json", fake_request)
        with pytest.raises(mfa_gate.GateError, match="fleet CHƯA hội tụ"):
            mfa_gate._probe_fleet(phien, "http://localhost:8000")

    def test_probe_fleet_thieu_luot_la_do_rieng(self, mfa_gate, monkeypatch):
        """Đủ-số-lượt và toàn-200 là HAI phép riêng.

        Gộp lại thì khi đỏ không biết đỏ vì thiếu lượt hay vì có mã lạ.
        """
        phien = mfa_gate.PhienDaXacThuc("admin", object(), object(), {}, None)
        goc = mfa_gate._probe_fleet.__globals__["PROBE_FLEET_LUOT"]
        assert goc == mfa_gate.PROBE_FLEET_LUOT
        nguon = inspect.getsource(mfa_gate._probe_fleet)
        assert nguon.count("raise GateError") == 2, (
            "cần ĐÚNG hai nhánh đỏ riêng: thiếu lượt, và mã khác 200"
        )

    # ------------------------------------------------------------------
    # Một nguồn chuẩn cho đường xác thực.
    # ------------------------------------------------------------------

    def test_sync_va_preflight_dung_chung_ham_dang_nhap(self, mfa_gate):
        """Cả hai đường phải đi qua `_dang_nhap_va_chung_minh`, không tự đăng nhập.

        Hai bản sao thì mọi hàng rào ở chặng chỉ-mật-khẩu phải vá hai lần, và
        lần thứ hai là lần bị quên.
        """
        fn = _ham_theo_ten(("sync_casbin", "_preflight_account"))
        for ten, node in fn.items():
            goi = {
                n.func.id
                for n in ast.walk(node)
                if isinstance(n, ast.Call) and isinstance(n.func, ast.Name)
            }
            assert "_dang_nhap_va_chung_minh" in goi, (
                f"{ten} không đi qua hàm đăng nhập chung"
            )
            chuoi = ast.dump(node)
            assert "/api/auth/login" not in chuoi, f"{ten} tự gọi /api/auth/login"
            assert "/api/auth/verify-mfa" not in chuoi, f"{ten} tự gọi verify-mfa"

    def test_preflight_account_khong_con_before_probe(self, mfa_gate):
        """Chữ ký `_preflight_account` không được còn tham số callback nào."""
        chu_ky = inspect.signature(mfa_gate._preflight_account)
        assert list(chu_ky.parameters) == ["account", "base_url"], chu_ky

        # Soi AST chứ không dò chuỗi: docstring CÓ nhắc tên cũ để giải thích vì
        # sao nó bị gỡ, và một phép `in` trên mã nguồn sẽ bắt trúng lời giải
        # thích ấy thay vì bắt mã.
        node = _ham_theo_ten(("_preflight_account",))["_preflight_account"]
        ten_dung = {
            n.id for n in ast.walk(node) if isinstance(n, ast.Name)
        } | {a.arg for a in ast.walk(node) if isinstance(a, ast.arg)}
        assert "before_probe" not in ten_dung, "mã vẫn dùng `before_probe`"

    def test_main_khai_du_bon_lenh(self, mfa_gate):
        """`sync-casbin` phải là lệnh THẬT, không phải nhánh chết."""
        nguon = inspect.getsource(mfa_gate.main)
        assert '"generate", "bootstrap", "sync-casbin", "preflight"' in nguon
        assert "sync_casbin()" in nguon

    def test_preflight_kiem_dung_vai_tra_ve(self, mfa_gate, monkeypatch):
        """Đăng nhập được KHÔNG chứng minh đúng quyền.

        Một cặp credential hợp lệ của vai khác vẫn cho 200 ở cả login lẫn
        verify; chỉ phép so `actual_role != expected_role` mới bắt được.
        """
        su_kien: list[str] = []
        responses = [
            (200, {"mfa_required": True, "mfa_token": "challenge"}),
            (401, {"detail": "Not authenticated"}),
            (200, {"user": {"role": "officer"}}),  # <-- sai vai
            (200, []),                             # KHÔNG được dùng tới
        ]

        def fake_request(opener, method, url, **kwargs):
            su_kien.append(method)
            return responses.pop(0)

        monkeypatch.setattr(mfa_gate, "_request_json", fake_request)
        monkeypatch.setattr(
            mfa_gate, "_totp_for_preflight", lambda *args: ("123456", 55_000_000)
        )
        with pytest.raises(mfa_gate.GateError) as loi:
            mfa_gate._preflight_account(
                (
                    "admin",
                    "admin",
                    "password",
                    "admin",
                    "JBSWY3DPEHPK3PXPJBSWY3DPEHPK3PXP",
                ),
                "http://localhost:8000",
            )
        assert "authenticated role is" in str(loi.value)
        assert "'officer'" in str(loi.value)

        assert su_kien == ["POST", "GET", "POST"], (
            "vai sai mà cổng vẫn đi tiếp (probe cuối phải KHÔNG được phát): %r"
            % (su_kien,)
        )
